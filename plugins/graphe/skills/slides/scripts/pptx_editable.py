"""An editable .pptx from a web deck, laid out by Chromium. Used by `slides.py pptx --editable`.

Each slide is shown fully built, then read element by element where the browser placed it:
- text becomes native text boxes, one per block, with bold, italic, colour, size, links and
  list bullets kept per run; inline TeX (data-tex, KaTeX or MathJax source) becomes Unicode;
- tables become native tables, with no theme table style, so Google Slides does not paint
  them blue, and their cell fills, borders and alignment;
- figures (img, svg, canvas, video) and display equations become pictures, with alt text;
- everything else (backgrounds, cards, rules, icons inside text) is baked into one picture
  set as the slide background, taken with the text glyphs made transparent, so it can not
  be dragged by mistake; a background of one colour is set as a solid fill instead;
- the first heading goes in the title placeholder where it stands, and the notes carry over.

Text boxes get width slack: PowerPoint and Google Slides set most fonts a little wider than
Chromium, and an unwrapped single line is set not to wrap at all. Sources for the choices
(Microsoft ResearchStudio html2pptx GOTCHAS: colours read through a canvas, native bullets,
no-wrap short lines, equations rasterized because slide math renders blank outside
PowerPoint) are cited in SKILL.md.

The deck is reveal.js, or a custom deck: one `<section class="slide">` per slide, shown by
the class `active`, notes in `<template class="speaker-notes">` or `<aside class="notes">`.
"""

# ruff: noqa: RUF001  # the TeX tables map commands to Greek and math characters
from __future__ import annotations

import io
import math
import re
import unicodedata
from dataclasses import dataclass, field
from typing import Any

# PowerPoint's slide is 13.333 in wide: 12192000 EMU, 960 pt.
SLIDE_W_EMU = 12192000
SLIDE_W_PT = 960.0
# "No Style, No Grid": a table with no theme style for Google Slides to repaint.
NO_TABLE_STYLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
# Width slack for a wrapped text box, as a share of its width plus a fixed margin in slide
# pixels (1600 across): other renderers set most fonts a little wider than Chromium.
SLACK_SHARE, SLACK_PX = 0.06, 10.0
# A table grows by the slack no further right than this share of the slide's width.
TABLE_RIGHT = 0.95
# Figures and equations are captured at twice the slide's pixel size, so they stay sharp.
PICTURE_SCALE = 2

# ---------------------------------------------------------------------- TeX --

SUP = dict(zip("0123456789+-=()niTt", "⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻⁼⁽⁾ⁿⁱᵀᵗ", strict=True))
SUB = dict(zip("0123456789+-=()aehijklmnoprstuvx", "₀₁₂₃₄₅₆₇₈₉₊₋₌₍₎ₐₑₕᵢⱼₖₗₘₙₒₚᵣₛₜᵤᵥₓ", strict=True))
SYMBOLS = {
    "times": "×", "cdot": "·", "approx": "≈", "le": "≤", "leq": "≤", "ge": "≥", "geq": "≥",
    "sim": "∼", "neq": "≠", "ne": "≠", "pm": "±", "mp": "∓", "to": "→", "rightarrow": "→",
    "leftarrow": "←", "Rightarrow": "⇒", "infty": "∞", "in": "∈", "notin": "∉", "subset": "⊂",
    "subseteq": "⊆", "cup": "∪", "cap": "∩", "forall": "∀", "exists": "∃", "nabla": "∇",
    "partial": "∂", "sum": "∑", "prod": "∏", "int": "∫", "top": "ᵀ", "ldots": "…", "dots": "…",
    "cdots": "⋯", "alpha": "α", "beta": "β", "gamma": "γ", "Gamma": "Γ", "delta": "δ",
    "Delta": "Δ", "epsilon": "ε", "varepsilon": "ε", "zeta": "ζ", "eta": "η", "theta": "θ",
    "Theta": "Θ", "kappa": "κ", "lambda": "λ", "Lambda": "Λ", "mu": "μ", "nu": "ν", "xi": "ξ",
    "pi": "π", "Pi": "Π", "rho": "ρ", "sigma": "σ", "Sigma": "Σ", "tau": "τ", "phi": "φ",
    "varphi": "φ", "Phi": "Φ", "chi": "χ", "psi": "ψ", "Psi": "Ψ", "omega": "ω", "Omega": "Ω",
    "|": "‖", "lVert": "‖", "rVert": "‖", "Vert": "‖", "lvert": "|", "rvert": "|",
    "odot": "⊙", "otimes": "⊗", "oplus": "⊕", "propto": "∝", "equiv": "≡", "mid": "|",
    "langle": "⟨", "rangle": "⟩", "{": "{", "}": "}", "%": "%", "$": "$", "_": "_", "&": "&",
    ",": " ", ";": " ", ":": " ", "!": "", " ": " ", "quad": " ", "qquad": "  ",
}  # fmt: skip
# Commands that wrap one argument and print it as it is.
_KEEP_ARGUMENT = {
    "mathrm", "mathbf", "boldsymbol", "text", "textrm", "textbf", "textit", "mathit", "mathsf",
    "mathcal", "mathbb", "mathtt", "operatorname", "bm", "underbrace", "overbrace", "overline",
    "underline", "widehat", "widetilde", "vec", "hat", "bar", "tilde",
}  # fmt: skip
# Accents with a Unicode combining mark, written after the argument.
_COMBINING = {"dot": "̇", "ddot": "̈"}
_SIZING = {
    "left", "right", "big", "Big", "bigg", "Bigg", "bigl", "bigr", "Bigl", "Bigr", "middle",
    "displaystyle", "textstyle", "scriptstyle", "limits", "nolimits",
}  # fmt: skip
_COMMAND = re.compile(r"\\([A-Za-z]+|.)")


def _group(s: str, i: int) -> tuple[str, int]:
    """(content, next index) of the {group} or single token at s[i]."""
    while i < len(s) and s[i] == " ":
        i += 1
    if i < len(s) and s[i] == "{":
        depth, j = 1, i + 1
        while j < len(s) and depth:
            depth += {"{": 1, "}": -1}.get(s[j], 0)
            j += 1
        return s[i + 1 : j - 1], j
    if i < len(s) and s[i] == "\\":
        m = _COMMAND.match(s, i)
        if m:
            return m.group(0), m.end()
    return s[i : i + 1], i + 1


_SCRIPTS = "".join(SUP.values()) + "".join(SUB.values())
# A number, or one symbol with its sub- and superscripts, needs no parentheses in a fraction.
_ATOM = re.compile(rf"\d+(?:\.\d+)?|[^\W\d_][{_SCRIPTS}]*")


def _wrap(x: str) -> str:
    return x if _ATOM.fullmatch(x) else f"({x})"


def tex(s: str) -> str:
    """Readable Unicode for inline TeX: Greek, operators, sub- and superscripts where Unicode
    has the characters (x_(ij) where it does not), fractions as a/b."""
    out: list[str] = []
    i = 0
    while i < len(s):
        c = s[i]
        if c in "^_":
            g, i = _group(s, i + 1)
            g = tex(g)
            table = SUP if c == "^" else SUB
            if g and all(ch in table.values() for ch in g):  # already raised, as \top is
                out.append(g)
            elif g and all(ch in table for ch in g):
                out.append("".join(table[ch] for ch in g))
            else:
                out.append(c + (g if len(g) == 1 else f"({g})"))
        elif c == "\\":
            m = _COMMAND.match(s, i)
            if not m:
                out.append(c)
                i += 1
                continue
            name, i = m.group(1), m.end()
            if name in ("frac", "dfrac", "tfrac"):
                a, i = _group(s, i)
                b, i = _group(s, i)
                out.append(f"{_wrap(tex(a))}/{_wrap(tex(b))}")
            elif name == "sqrt":
                a, i = _group(s, i)
                out.append(f"√{_wrap(tex(a))}")
            elif name in _KEEP_ARGUMENT:
                a, i = _group(s, i)
                out.append(tex(a))
            elif name in _COMBINING:
                a, i = _group(s, i)
                out.append(tex(a) + _COMBINING[name])
            elif name in _SIZING:
                pass
            else:
                out.append(SYMBOLS.get(name, name))
        elif c in "{}":
            i += 1
        elif c == "~":
            out.append(" ")
            i += 1
        else:
            out.append(c)
            i += 1
    # Composed form: an accent written after its letter (ẍ) becomes the one character.
    return unicodedata.normalize("NFC", re.sub(r" {2,}", " ", "".join(out)))


# ---------------------------------------------------------------- browser --

# The custom deck contract: section.slide, the one on screen marked `active`.
CUSTOM_SLIDES = """() => {
  window.__deckSlides = () => [...document.querySelectorAll('section.slide')];
  // What is not the slide (the deck's own buttons, a counter) stays out of the export.
  const iso = document.createElement('style');
  iso.textContent = '[data-pptx-outside] { visibility: hidden !important; }';
  document.head.appendChild(iso);
  window.__deckShow = (i) => {
    const out = 'data-pptx-outside';
    document.querySelectorAll(`[${out}]`).forEach((el) => el.removeAttribute(out));
    let a = __deckSlides()[i];
    for (; a && a.parentElement && a !== document.body; a = a.parentElement) {
      for (const sib of a.parentElement.children) if (sib !== a) sib.setAttribute(out, '');
    }
    __deckSlides().forEach((s, k) => {
      const on = k === i;
      s.classList.toggle('active', on);
      if (on) ['inert', 'hidden', 'aria-hidden'].forEach((a) => s.removeAttribute(a));
    });
    for (const f of __deckSlides()[i].querySelectorAll('.fragment')) f.classList.add('visible');
  };
  window.__deckFrame = (s) => {
    const r = s.getBoundingClientRect();
    return {left: r.left, top: r.top, width: r.width, height: r.height,
            scale: r.width / (s.offsetWidth || r.width)};
  };
  window.__deckNotes = (s) => {
    const n = s.querySelector('template.speaker-notes, aside.notes');
    return n ? n.innerHTML : '';
  };
}"""

# reveal.js: the frame is the slide canvas (config size times scale), as _READ_SLIDE reads it.
REVEAL_SLIDES = """() => {
  window.__deckShow = (i) => {
    const s = __deckSlides()[i];
    const at = Reveal.getIndices(s);
    Reveal.slide(at.h, at.v);
    while (Reveal.nextFragment()) {}
    for (const f of s.querySelectorAll('.fragment')) f.classList.add('visible');
  };
  window.__deckFrame = () => {
    const config = Reveal.getConfig(), scale = Reveal.getScale();
    const r = document.querySelector('.reveal .slides').getBoundingClientRect();
    return {left: r.left, top: r.top, width: config.width * scale,
            height: config.height * scale, scale};
  };
  window.__deckNotes = (s) => Reveal.getSlideNotes(s) || '';
}"""

# Read the shown slide i: its frame, notes, and every text block, table, figure and display
# equation in document order, in screen pixels relative to the frame. Each figure and
# equation is tagged with data-pptx-shot so Python can capture it.
READ = r"""(i) => {
  const s = __deckSlides()[i];
  const f = __deckFrame(s);
  const canvas = document.createElement('canvas');
  canvas.width = canvas.height = 1;
  const ctx = canvas.getContext('2d', {willReadFrequently: true});
  // Any CSS colour (oklab, color-mix, named) as [r, g, b, alpha], resolved by the browser.
  const rgba = (c) => {
    if (!c || c === 'transparent') return [0, 0, 0, 0];
    ctx.clearRect(0, 0, 1, 1);
    ctx.fillStyle = 'rgba(0, 0, 0, 0)';
    ctx.fillStyle = c;
    ctx.fillRect(0, 0, 1, 1);
    const d = ctx.getImageData(0, 0, 1, 1).data;
    return [d[0], d[1], d[2], +(d[3] / 255).toFixed(2)];
  };
  const BLOCK = new Set(['ADDRESS', 'ARTICLE', 'ASIDE', 'BLOCKQUOTE', 'CAPTION', 'DD', 'DETAILS',
    'DIALOG', 'DIV', 'DL', 'DT', 'FIELDSET', 'FIGCAPTION', 'FIGURE', 'FOOTER', 'FORM', 'H1', 'H2',
    'H3', 'H4', 'H5', 'H6', 'HEADER', 'HGROUP', 'HR', 'LI', 'MAIN', 'NAV', 'OL', 'P', 'PRE',
    'SECTION', 'SUMMARY', 'TABLE', 'UL', 'TBODY', 'THEAD', 'TFOOT', 'TR', 'TD', 'TH']);
  const MEDIA = new Set(['IMG', 'SVG', 'CANVAS', 'VIDEO', 'PICTURE', 'IFRAME', 'OBJECT', 'EMBED']);
  const SKIP = new Set(['SCRIPT', 'STYLE', 'TEMPLATE', 'NOSCRIPT', 'BUTTON', 'INPUT', 'SELECT',
    'TEXTAREA']);
  const tag = (el) => el.tagName.toUpperCase();
  const px = (r) => ({x: r.left - f.left, y: r.top - f.top, w: r.width, h: r.height});
  const inFrame = (r) => r.width >= 1 && r.height >= 1 && r.right > f.left
    && r.bottom > f.top && r.left < f.left + f.width && r.top < f.top + f.height;
  // An element with display: contents has no box of its own; its children show for it.
  const shown = (el) => (getComputedStyle(el).display === 'contents'
    ? getComputedStyle(el).visibility !== 'hidden' && (!el.parentElement || shown(el.parentElement))
    : el.checkVisibility({opacityProperty: true, visibilityProperty: true}));
  // Screen-reader-only text: inside a box of 1px or less that clips what it holds, measured
  // in the page's own pixels, since the deck may be scaled on screen.
  const srOnly = (el) => {
    for (let a = el; a && a !== s; a = a.parentElement) {
      const st = getComputedStyle(a), r = a.getBoundingClientRect();
      const clips = st.overflow !== 'visible' || (st.clip && st.clip !== 'auto')
        || st.clipPath !== 'none';
      const w = a.offsetWidth ?? r.width / f.scale, h = a.offsetHeight ?? r.height / f.scale;
      if (w <= 1 && h <= 1 && clips) return true;
    }
    return false;
  };
  const skipped = (el) => SKIP.has(tag(el)) || el.matches('aside.notes, .slide-number')
    || !shown(el) || srOnly(el);
  const texSource = (el) => {
    if (el.dataset && el.dataset.tex != null) return el.dataset.tex;
    const a = el.querySelector('annotation[encoding="application/x-tex"]');
    if (a) return a.textContent;
    try {
      const item = window.MathJax && MathJax.startup.document.getMathItemsWithin(el)[0];
      if (item) return item.math;
    } catch (e) {}
    return null;
  };
  const DISPLAY_MATH = '.equation, .katex-display, mjx-container[display="true"], '
    + 'math[display="block"]';
  const displayMath = (el) => el.matches(DISPLAY_MATH)
    || (el.dataset.tex != null && getComputedStyle(el).display === 'block');
  const inlineMath = (el) => !displayMath(el) && (el.dataset.tex != null
    || el.matches('.katex, mjx-container, math'));
  const hasBlock = (el) => [...el.querySelectorAll('*')].some((d) => BLOCK.has(tag(d))
    && !d.closest('.katex, mjx-container, math, svg'));
  const placed = (el) => ['absolute', 'fixed'].includes(getComputedStyle(el).position);
  // An inline tag set to display as a block (a stat's number over its label, or any flex or
  // grid item, which CSS blockifies) gets a box of its own; its parent keeps the rest.
  const blocky = (el) => /^(block|flex|grid|list-item|table|flow-root)$/
    .test(getComputedStyle(el).display);
  const items = [];
  let shot = 0;

  // An equation is cut to its content: a display block spans the line, its formula does not.
  const tight = (el) => {
    const range = document.createRange();
    range.selectNodeContents(el);
    const c = range.getBoundingClientRect();
    return c.width >= 1 && c.height >= 1 ? c : el.getBoundingClientRect();
  };
  const picture = (el, kind, alt) => {
    const r = kind === 'equation' ? tight(el) : el.getBoundingClientRect();
    if (!inFrame(r)) return;
    el.setAttribute('data-pptx-shot', String(shot));
    items.push({kind, shot: shot++, ...px(r), alt: (alt || '').trim(),
                box: {x: r.left, y: r.top, width: r.width, height: r.height}});
  };
  const figureAlt = (el) => el.getAttribute('alt') || el.getAttribute('aria-label')
    || (el.querySelector && el.querySelector('title') && el.querySelector('title').textContent)
    || '';

  // Generated text: quoted strings and attr() become a run. Content the page cannot hand
  // over (a counter, a quote mark, an image) stays in the background, drawn in its own colour.
  const keep = [];
  const unescape = (t) => t.replace(/\\([0-9a-fA-F]{1,6}) ?/g,
    (_, h) => String.fromCodePoint(parseInt(h, 16))).replace(/\\(.)/g, '$1');
  const pseudo = (el, which) => {
    const st = getComputedStyle(el, which), c = st.content;
    if (!c || c === 'none' || c === 'normal') return '';
    let out = '', whole = true;
    const token = /"((?:[^"\\]|\\.)*)"|attr\(\s*([\w-]+)\s*\)|(\S+\([^)]*\)|\S+)/gs;
    for (const m of c.matchAll(token)) {
      if (m[1] !== undefined) out += unescape(m[1]);
      else if (m[2]) out += el.getAttribute(m[2]) || '';
      else whole = false;
    }
    if (whole) return out;
    keep.push({el, which, color: st.color});
    return '';
  };
  const runStyle = (el) => {
    const st = getComputedStyle(el);
    const link = el.closest('a[href]');
    return {
      size: parseFloat(st.fontSize) * f.scale,
      bold: (parseInt(st.fontWeight, 10) || 400) >= 600,
      italic: st.fontStyle !== 'normal',
      underline: st.textDecorationLine.includes('underline'),
      strike: st.textDecorationLine.includes('line-through'),
      color: rgba(st.color),
      family: st.fontFamily.split(',')[0].trim().replace(/^["']|["']$/g, ''),
      shift: st.verticalAlign === 'super' ? 30 : st.verticalAlign === 'sub' ? -25 : 0,
      spacing: st.letterSpacing === 'normal' ? 0 : (parseFloat(st.letterSpacing) || 0) * f.scale,
      href: link && s.contains(link) ? link.href : null,
      transform: st.textTransform,
      pre: /^pre/.test(st.whiteSpace),
    };
  };
  const cased = (text, how) => how === 'uppercase' ? text.toUpperCase()
    : how === 'lowercase' ? text.toLowerCase()
    : how === 'capitalize' ? text.replace(/\b\w/g, (c) => c.toUpperCase()) : text;
  // Runs of an inline group: text nodes with their element's style, <br> as a line break,
  // inline math as its TeX source, small media left to the background.
  const glyphs = new Set(), math = new Set(), listed = new Set();
  const runsOf = (nodes) => {
    const runs = [];
    const visit = (n) => {
      if (n.nodeType === 3) {
        const el = n.parentElement;
        if (!el || !shown(el) || !n.textContent) return;
        // Text the browser paints through its background (background-clip: text) or not at
        // all has no colour of its own; it stays in the background picture.
        const css = getComputedStyle(el);
        if (rgba(css.webkitTextFillColor)[3] === 0 || rgba(css.color)[3] === 0) return;
        glyphs.add(el);
        const st = runStyle(el);
        runs.push({...st, text: cased(n.textContent, st.transform)});
        return;
      }
      if (n.nodeType !== 1 || skipped(n)) return;
      if (tag(n) === 'BR') {
        runs.push({...runStyle(n.parentElement), text: '\n', pre: true});
        return;
      }
      if (MEDIA.has(tag(n))) return;
      if (inlineMath(n)) {
        const source = texSource(n);
        math.add(n);
        if (source != null) runs.push({...runStyle(n), tex: source, text: ''});
        else runs.push({...runStyle(n), text: n.textContent});
        return;
      }
      const before = pseudo(n, '::before'), after = pseudo(n, '::after');
      if (before) { glyphs.add(n); runs.push({...runStyle(n), text: before}); }
      n.childNodes.forEach(visit);
      if (after) { glyphs.add(n); runs.push({...runStyle(n), text: after}); }
    };
    nodes.forEach(visit);
    return runs;
  };
  const listInfo = (li) => {
    const list = li.parentElement;
    const st = getComputedStyle(li);
    if (!list || !['UL', 'OL'].includes(tag(list)) || st.listStyleType === 'none') return null;
    const ordered = tag(list) === 'OL' || /decimal|alpha|roman/.test(st.listStyleType);
    let number = 1;
    if (ordered) {
      const siblings = [...list.children].filter((c) => tag(c) === 'LI');
      number = (list.start || 1) + siblings.indexOf(li);
      if (li.value) number = li.value;
    }
    // The marker sits in the room left of the item: the list's padding, or its margin where
    // a reset (reveal.js) zeroes the padding. At least 1em, so the bullet never touches text.
    const em = parseFloat(st.fontSize) * f.scale;
    const room = li.getBoundingClientRect().left - list.getBoundingClientRect().left
      + (parseFloat(getComputedStyle(list).marginLeft) || 0) * f.scale;
    const inside = st.listStylePosition === 'inside';
    return {ordered, number, type: st.listStyleType, color: rgba(st.color),
            indent: inside ? 0 : Math.min(Math.max(room, em), em * 1.6)};
  };
  // The list item a block opens: the item itself, or its first child when the item's
  // content starts with a block (<li><p>...</p></li>).
  const itemOf = (c, first) => {
    if (tag(c) === 'LI') return first ? c : null;
    const li = c.parentElement;
    if (!li || tag(li) !== 'LI') return null;
    for (const k of li.childNodes) {
      if (k === c) return li;
      if (k.nodeType === 3 ? k.textContent.trim() : k.nodeType === 1 && !skipped(k)) return null;
    }
    return null;
  };
  const textBlock = (container, nodes, first, last) => {
    const runs = runsOf(nodes);
    // The block's own generated text opens its first line and closes its last.
    const before = first ? pseudo(container, '::before') : '';
    const after = last ? pseudo(container, '::after') : '';
    if (before) runs.unshift({...runStyle(container), text: before});
    if (after) runs.push({...runStyle(container), text: after});
    if (before || after) glyphs.add(container);
    if (!runs.some((r) => r.tex != null || r.text.trim())) return;
    const range = document.createRange();
    range.setStartBefore(nodes[0]);
    range.setEndAfter(nodes[nodes.length - 1]);
    const r = range.getBoundingClientRect();
    if (!inFrame(r)) return;
    const st = getComputedStyle(container);
    const lh = parseFloat(st.lineHeight);
    // The range is the glyphs' box; a line box adds half the leading above and below, and
    // PowerPoint sets exact line spacing the same way, so the text box takes the line box.
    const lines = [...range.getClientRects()].filter((q) => q.width > 0);
    const tall = lines.length ? lines[0].height : r.height;
    const half = Number.isFinite(lh) ? Math.max(0, (lh * f.scale - tall) / 2) : 0;
    const box = {x: r.left - f.left, y: r.top - f.top - half, w: r.width, h: r.height + 2 * half};
    // Lines the browser set: a rect that starts below the last line's bottom opens a new one.
    let count = 0, bottom = -Infinity;
    for (const q of lines.sort((a, b) => a.top - b.top)) {
      if (q.top >= bottom - 1) {
        count += 1;
        bottom = q.bottom;
      } else bottom = Math.max(bottom, q.bottom);
    }
    // Wrapped text was set to the container's width, not its longest line's.
    if (count > 1) {
      const c = container.getBoundingClientRect();
      const inner = c.right - ((parseFloat(st.paddingRight) || 0)
        + (parseFloat(st.borderRightWidth) || 0)) * f.scale;
      box.w = Math.max(box.w, inner - r.left);
    }
    const item = itemOf(container, first);
    const li = item && !listed.has(item) ? listInfo(item) : null;
    if (li) { listed.add(item); glyphs.add(item); }
    items.push({
      kind: 'text', ...box, runs,
      tag: tag(container).toLowerCase(),
      align: st.textAlign,
      rtl: st.direction === 'rtl',
      lineHeight: Number.isFinite(lh) ? lh * f.scale : null,
      lines: count,
      fontSize: parseFloat(st.fontSize) * f.scale,
      bullet: li,
    });
  };
  const cellOf = (td) => {
    const r = td.getBoundingClientRect(), st = getComputedStyle(td);
    let bg = rgba(st.backgroundColor);
    for (let a = td.parentElement; bg[3] === 0 && a && a !== s; a = a.parentElement) {
      bg = rgba(getComputedStyle(a).backgroundColor);
      if (tag(a) === 'TABLE') break;
    }
    const side = (k) => ({width: (parseFloat(st[`border${k}Width`]) || 0) * f.scale,
      color: rgba(st[`border${k}Color`]), style: st[`border${k}Style`]});
    td.querySelectorAll('*').forEach((d) => glyphs.add(d));
    glyphs.add(td);
    return {
      ...px(r), runs: runsOf([...td.childNodes]), header: tag(td) === 'TH',
      colspan: td.colSpan || 1, rowspan: td.rowSpan || 1, bg,
      align: st.textAlign, valign: st.verticalAlign,
      pad: ['Top', 'Right', 'Bottom', 'Left']
        .map((k) => (parseFloat(st[`padding${k}`]) || 0) * f.scale),
      borders: {top: side('Top'), right: side('Right'), bottom: side('Bottom'), left: side('Left')},
      lineHeight: (parseFloat(st.lineHeight) || 0) * f.scale || null,
    };
  };
  const table = (el) => {
    const r = el.getBoundingClientRect();
    if (!inFrame(r)) return;
    el.setAttribute('data-pptx-table', '');
    const rows = [...el.rows].filter((tr) => shown(tr)).map((tr) => ({
      h: tr.getBoundingClientRect().height,
      cells: [...tr.cells].filter((td) => shown(td)).map(cellOf),
    }));
    if (rows.some((row) => row.cells.length)) items.push({kind: 'table', ...px(r), rows});
  };

  const walk = (el) => {
    let group = [];
    let first = true;
    const flush = (last = false) => {
      if (group.length) { textBlock(el, group, first, last); first = false; }
      group = [];
    };
    for (const n of el.childNodes) {
      if (n.nodeType === 3) { if (n.textContent) group.push(n); continue; }
      if (n.nodeType !== 1 || skipped(n)) continue;
      const t = tag(n);
      if (displayMath(n)) {
        flush();
        picture(n, 'equation', texSource(n) || n.textContent);
        continue;
      }
      if (t === 'TABLE') { flush(); table(n); continue; }
      if (MEDIA.has(t)) {
        const r = n.getBoundingClientRect();
        const line = parseFloat(getComputedStyle(el).fontSize) * f.scale * 2;
        const inline = group.some((g) => g.textContent && g.textContent.trim()) && r.height <= line;
        if (inline) { group.push(n); continue; }
        flush();
        picture(n, 'figure', figureAlt(n));
        continue;
      }
      if (inlineMath(n)) { group.push(n); continue; }
      if (BLOCK.has(t) || placed(n) || hasBlock(n) || blocky(n)) {
        flush();
        walk(n);
        continue;
      }
      group.push(n);
    }
    flush(true);
  };
  walk(s);

  const heading = items.find((b) => b.kind === 'text' && /^h[1-3]$/.test(b.tag));
  if (heading) heading.title = true;
  for (const el of glyphs) el.setAttribute('data-pptx-glyphs', '');
  for (const el of math) el.setAttribute('data-pptx-math', '');
  const rules = keep.map(({el, which, color}, k) => {
    el.setAttribute(`data-pptx-keep-${which.slice(2)}`, String(k));
    return `[data-pptx-keep-${which.slice(2)}="${k}"]${which} { color: ${color} !important;`
      + ` -webkit-text-fill-color: ${color} !important; }`;
  });
  return {frame: f, items, notesHtml: __deckNotes(s), keep: rules.join('\n')};
}"""

# Text made transparent, so figures and the background are taken without it. Only the glyph
# fill goes: `color` stays, so an icon or border drawn in currentColor keeps its colour.
# Inline math is drawn with borders and SVG in its colour, so it loses that colour too.
# Figures and equations get their glyphs back for their own pictures.
GLYPH_STYLE = """
[data-pptx-glyphs], [data-pptx-glyphs]::before, [data-pptx-glyphs]::after,
[data-pptx-glyphs]::first-letter, [data-pptx-glyphs]::first-line {
  -webkit-text-fill-color: transparent !important; text-shadow: none !important;
  text-decoration-color: transparent !important;
}
[data-pptx-glyphs]::marker, [data-pptx-math], [data-pptx-math] * {
  color: transparent !important; -webkit-text-fill-color: transparent !important;
}
[data-pptx-shot], [data-pptx-shot] * { -webkit-text-fill-color: currentcolor !important; }"""
# The background pass also hides the tables, figures and equations placed on their own.
PLACED_STYLE = """
[data-pptx-table], [data-pptx-table] *, [data-pptx-shot], [data-pptx-shot] * {
  visibility: hidden !important;
}"""
STYLE = """([id, css]) => {
  const st = document.createElement('style');
  st.id = id;
  st.textContent = css;
  document.head.appendChild(st);
}"""
# Backgrounds cleared under the equations, for a capture with only the formula's own pixels.
CLEAR_STYLE = """
[data-pptx-clear] { background: transparent !important; box-shadow: none !important; }
.reveal .backgrounds { visibility: hidden !important; }"""
CLEAR = """() => {
  for (const el of document.querySelectorAll('[data-pptx-shot]')) {
    for (let a = el; a; a = a.parentElement) a.setAttribute('data-pptx-clear', '');
  }
}"""
DROP = """(ids) => {
  for (const id of ids) {
    const st = document.getElementById(id);
    if (st) st.remove();
  }
}"""
UNMARK = """(ids) => {
  for (const id of ids) {
    const st = document.getElementById(id);
    if (st) st.remove();
  }
  const marks = ['data-pptx-glyphs', 'data-pptx-math', 'data-pptx-shot', 'data-pptx-table',
                 'data-pptx-clear', 'data-pptx-keep-before', 'data-pptx-keep-after'];
  for (const a of marks) {
    document.querySelectorAll(`[${a}]`).forEach((el) => el.removeAttribute(a));
  }
}"""
ALL_STYLES = ["pptx-glyphs", "pptx-clear", "pptx-placed"]
# A note's text with its paragraph breaks, laid out off-screen so innerText can read it.
NOTE_TEXT = """(html) => {
  const scratch = document.createElement('div');
  scratch.style.cssText = 'position:absolute;left:-99999px;top:0;width:800px';
  scratch.innerHTML = html;
  scratch.querySelectorAll('template').forEach((t) => t.replaceWith(t.content));
  document.body.appendChild(scratch);
  const text = scratch.innerText.trim().replace(/\\n\\s*\\n/g, '\\n');
  scratch.remove();
  return text;
}"""


@dataclass
class Captured:
    """One slide as read and photographed in the browser."""

    frame: dict
    items: list
    notes: str
    background: bytes
    shots: dict = field(default_factory=dict)
    image: bytes | None = None


def is_custom(page) -> bool:
    """A page with section.slide and no reveal.js: the custom deck contract."""
    return page.evaluate(
        "!document.querySelector('.reveal .slides') && "
        "document.querySelectorAll('section.slide').length > 0"
    )


def _clip(box: dict, frame: dict) -> dict | None:
    """box cut to the frame, as a Playwright screenshot clip, or None when nothing is left."""
    x0, y0 = max(box["x"], frame["left"]), max(box["y"], frame["top"])
    x1 = min(box["x"] + box["width"], frame["left"] + frame["width"])
    y1 = min(box["y"] + box["height"], frame["top"] + frame["height"])
    if x1 - x0 < 1 or y1 - y0 < 1:
        return None
    return {"x": x0, "y": y0, "width": x1 - x0, "height": y1 - y0}


def _read(page, i: int) -> dict:
    """Show slide i and read it, with the window grown first when the slide is larger."""
    for _ in range(2):
        page.evaluate("(i) => __deckShow(i)", i)
        page.evaluate(
            "() => new Promise((r) => requestAnimationFrame(() => requestAnimationFrame(r)))"
        )
        read = page.evaluate(READ, i)
        f, view = read["frame"], page.viewport_size
        need = (math.ceil(f["left"] + f["width"]), math.ceil(f["top"] + f["height"]))
        if need[0] <= view["width"] and need[1] <= view["height"]:
            return read
        # A screenshot stops at the window's edge, so the window grows to hold the slide.
        page.evaluate(UNMARK, ALL_STYLES)
        page.set_viewport_size(
            {"width": max(view["width"], need[0]), "height": max(view["height"], need[1])}
        )
    return read


def capture(page, count: int, still: str, keep_images: bool = False) -> list[Captured]:
    """Show, read and photograph each slide. `still` is the style that stops motion."""
    page.add_style_tag(content=still)
    page.evaluate("document.fonts.ready.then(() => true)")
    out = []
    for i in range(count):
        read = _read(page, i)
        frame = read["frame"]
        whole = _clip({"x": frame["left"], "y": frame["top"], "width": frame["width"],
                       "height": frame["height"]}, frame)  # fmt: skip
        if whole is None:
            raise ValueError(f"slide {i + 1} has no area on screen")
        image = page.screenshot(clip=whole, scale="css") if keep_images else None
        # Text goes transparent first, so text laid over a figure is not in its picture.
        page.evaluate(STYLE, ["pptx-glyphs", GLYPH_STYLE + "\n" + read["keep"]])
        shots = {}
        for kind in ("figure", "equation"):
            if kind == "equation":  # taken without what lies behind it, to sit on any fill
                page.evaluate(CLEAR)
                page.evaluate(STYLE, ["pptx-clear", CLEAR_STYLE])
            for item in read["items"]:
                if item.get("kind") != kind:
                    continue
                clip = _clip(item["box"], frame)
                if clip is None:
                    continue
                shots[item["shot"]] = page.screenshot(clip=clip, omit_background=kind == "equation")
                item.update(x=clip["x"] - frame["left"], y=clip["y"] - frame["top"],
                            w=clip["width"], h=clip["height"])  # fmt: skip
        page.evaluate(DROP, ["pptx-clear"])
        page.evaluate(STYLE, ["pptx-placed", PLACED_STYLE])
        background = page.screenshot(clip=whole, scale="css")
        page.evaluate(UNMARK, ALL_STYLES)
        notes = page.evaluate(NOTE_TEXT, read["notesHtml"]) if read["notesHtml"] else ""
        out.append(Captured(frame, read["items"], notes, background, shots, image))
    return out


# ------------------------------------------------------------------- build --


def _hex(c) -> str:
    return "".join(f"{v:02X}" for v in c[:3])


def _alpha(srgb, alpha: float) -> None:
    """Set a colour element's opacity (DrawingML counts it in thousandths of a percent)."""
    from lxml import etree  # ty: ignore[unresolved-import]
    from pptx.oxml.ns import qn  # ty: ignore[unresolved-import]

    etree.SubElement(srgb, qn("a:alpha"), {"val": str(round(alpha * 100000))})


def rpr_colour(run):
    from pptx.oxml.ns import qn  # ty: ignore[unresolved-import]

    return run._r.get_or_add_rPr().find(f"{qn('a:solidFill')}/{qn('a:srgbClr')}")


def _rgb(c):
    from pptx.dml.color import RGBColor  # ty: ignore[unresolved-import]

    return RGBColor(*c[:3])


def uniform_colour(png: bytes) -> tuple[int, int, int] | None:
    """The one colour a picture is filled with, or None."""
    from PIL import Image  # ty: ignore[unresolved-import]

    with Image.open(io.BytesIO(png)) as im:
        rgb = im.convert("RGB")
        lo_hi = rgb.getextrema()
        if all(lo == hi for lo, hi in lo_hi):
            return tuple(lo for lo, _ in lo_hi)
    return None


def _set_background_picture(slide, png: bytes) -> None:
    """Put png in the slide's background, where it cannot be selected or dragged."""
    from lxml import etree  # ty: ignore[unresolved-import]
    from pptx.oxml.ns import qn  # ty: ignore[unresolved-import]

    _, rid = slide.part.get_or_add_image_part(io.BytesIO(png))
    c_sld = slide._element.find(qn("p:cSld"))
    old = c_sld.find(qn("p:bg"))
    if old is not None:
        c_sld.remove(old)
    bg = etree.Element(qn("p:bg"))
    c_sld.insert(0, bg)  # the schema puts the background first
    bg_pr = etree.SubElement(bg, qn("p:bgPr"))
    blip_fill = etree.SubElement(bg_pr, qn("a:blipFill"), {"dpi": "0", "rotWithShape": "1"})
    etree.SubElement(blip_fill, qn("a:blip"), {qn("r:embed"): rid})
    etree.SubElement(blip_fill, qn("a:srcRect"))
    stretch = etree.SubElement(blip_fill, qn("a:stretch"))
    etree.SubElement(stretch, qn("a:fillRect"))
    etree.SubElement(bg_pr, qn("a:effectLst"))


# CSS generic families have no font of that name in PowerPoint: the usual concrete stand-ins.
GENERIC_FAMILIES = {
    "sans-serif": "Arial",
    "system-ui": "Arial",
    "ui-sans-serif": "Arial",
    "-apple-system": "Arial",
    "serif": "Times New Roman",
    "ui-serif": "Times New Roman",
    "monospace": "Courier New",
    "ui-monospace": "Courier New",
}
_ALIGN = {"center": "CENTER", "right": "RIGHT", "end": "RIGHT", "justify": "JUSTIFY"}
_BULLET = {
    "disc": "•",
    "circle": "◦",
    "square": "▪",
    "disclosure-closed": "▸",
    "disclosure-open": "▾",
}
_NUMBERING = {
    "lower-alpha": "alphaLcPeriod",
    "lower-latin": "alphaLcPeriod",
    "upper-alpha": "alphaUcPeriod",
    "upper-latin": "alphaUcPeriod",
    "lower-roman": "romanLcPeriod",
    "upper-roman": "romanUcPeriod",
}


def _alignment(css: str, rtl: bool):
    from pptx.enum.text import PP_ALIGN  # ty: ignore[unresolved-import]

    if css == "start":
        css = "right" if rtl else "left"
    return getattr(PP_ALIGN, _ALIGN.get(css, "LEFT"))


_RUN_KEYS = (
    "size", "bold", "italic", "underline", "strike", "color", "family", "shift", "href", "spacing"
)  # fmt: skip


def _collapse(runs: list[dict]) -> list[dict]:
    """Runs with whitespace collapsed as the browser does outside `pre`, inline TeX turned
    to Unicode, neighbours of one style joined, and the block's ends trimmed."""
    out: list[dict] = []
    for r in runs:
        text = tex(r["tex"]) if r.get("tex") is not None else r["text"]
        if not r.get("pre"):
            text = re.sub(r"[ \t\r\n\f]+", " ", text)
        if out and out[-1]["text"].endswith((" ", "\n")) and text.startswith(" "):
            text = text[1:]
        if not text:
            continue
        if out and all(out[-1].get(k) == r.get(k) for k in _RUN_KEYS):
            out[-1]["text"] += text
        else:
            out.append({**r, "text": text})
    if out:
        out[0]["text"] = out[0]["text"].lstrip(" ")
        out[-1]["text"] = out[-1]["text"].rstrip(" ")
    for r in out:
        if not r.get("pre"):
            r["text"] = re.sub(r" ?\n ?", "\n", r["text"])
    return [r for r in out if r["text"]]


def _fill_text(tf, runs: list[dict], pt_per_px: float, block: dict, fonts: set[str]) -> None:
    """Write runs into a text frame: one paragraph, a line break per newline."""
    from pptx.util import Pt  # ty: ignore[unresolved-import]

    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = _alignment(block.get("align", "left"), block.get("rtl", False))
    if block.get("rtl"):
        p._p.get_or_add_pPr().set("rtl", "1")
    if block.get("lineHeight"):
        p.line_spacing = Pt(round(block["lineHeight"] * pt_per_px * 2) / 2)
    for r in _collapse(runs):
        parts = r["text"].split("\n")
        for k, part in enumerate(parts):
            if k:
                p.add_line_break()
            if not part:
                continue
            run = p.add_run()
            run.text = part
            font = run.font
            font.size = Pt(max(1.0, round(r["size"] * pt_per_px * 2) / 2))
            font.bold = r["bold"]
            font.italic = r["italic"]
            font.underline = r["underline"]
            family = GENERIC_FAMILIES.get(r["family"].lower(), r["family"])
            if family:
                font.name = family
                fonts.add(family)
            if r["color"] and r["color"][3] > 0:
                font.color.rgb = _rgb(r["color"])
                if r["color"][3] < 1:
                    _alpha(rpr_colour(run), r["color"][3])
            rpr = run._r.get_or_add_rPr()
            if r["strike"]:
                rpr.set("strike", "sngStrike")
            if r["shift"]:
                rpr.set("baseline", f"{r['shift'] * 1000}")
            if r.get("spacing"):
                rpr.set("spc", str(round(r["spacing"] * pt_per_px * 100)))
            if r.get("href") and re.match(r"(https?|mailto):", r["href"]):
                run.hyperlink.address = r["href"]


def _bullet(p, info: dict, pt_per_px: float, emu_per_px: float) -> None:
    """A native bullet or number with a hanging indent, as PowerPoint draws its own lists."""
    from lxml import etree  # ty: ignore[unresolved-import]
    from pptx.oxml.ns import qn  # ty: ignore[unresolved-import]

    ppr = p._p.get_or_add_pPr()
    indent = int(info["indent"] * emu_per_px)
    ppr.set("marL", str(indent))
    ppr.set("indent", str(-indent))
    if info["color"][3] > 0:
        clr = etree.SubElement(ppr, qn("a:buClr"))
        etree.SubElement(clr, qn("a:srgbClr"), {"val": _hex(info["color"])})
    if info["ordered"]:
        scheme = _NUMBERING.get(info["type"], "arabicPeriod")
        etree.SubElement(ppr, qn("a:buAutoNum"), {"type": scheme, "startAt": str(info["number"])})
    else:
        etree.SubElement(ppr, qn("a:buFont"), {"typeface": "Arial"})
        etree.SubElement(ppr, qn("a:buChar"), {"char": _BULLET.get(info["type"], "•")})


def _text_box(shapes, item: dict, geo: Geometry, fonts: set[str], placeholder=None):
    from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE  # ty: ignore[unresolved-import]

    x, y, w, h = item["x"], item["y"], item["w"], item["h"]
    bullet = item.get("bullet")
    if bullet:
        x -= bullet["indent"]
        w += bullet["indent"]
    one_line = item.get("lines") == 1
    if not one_line:
        w = max(w, min(w * (1 + SLACK_SHARE) + SLACK_PX * geo.px_per_slide_px, geo.width - x))
    if placeholder is not None:
        shape = placeholder
        shape.left, shape.top = geo.emu(x), geo.emu(y)
        shape.width, shape.height = geo.emu(w), geo.emu(h)
    else:
        shape = shapes.add_textbox(geo.emu(x), geo.emu(y), geo.emu(w), geo.emu(h))
    tf = shape.text_frame
    tf.word_wrap = not one_line
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    _fill_text(tf, item["runs"], geo.pt_per_px, item, fonts)
    if bullet:
        _bullet(tf.paragraphs[0], bullet, geo.pt_per_px, geo.emu_per_px)
    return shape


def _cell_border(cell, side: str, border: dict, geo: Geometry) -> None:
    from lxml import etree  # ty: ignore[unresolved-import]
    from pptx.oxml.ns import qn  # ty: ignore[unresolved-import]

    tag = {"left": "a:lnL", "right": "a:lnR", "top": "a:lnT", "bottom": "a:lnB"}[side]
    tc_pr = cell._tc.get_or_add_tcPr()
    old = tc_pr.find(qn(tag))
    if old is not None:
        tc_pr.remove(old)
    visible = (
        border["width"] > 0 and border["style"] not in ("none", "hidden") and border["color"][3] > 0
    )
    # Hairlines stay at least one rendered pixel (GOTCHAS G20).
    width = max(int(border["width"] * geo.emu_per_px), 12700) if visible else 0
    ln = etree.Element(qn(tag), {"w": str(width)})
    if visible:
        fill = etree.SubElement(ln, qn("a:solidFill"))
        hexa = _hex(border["color"])
        etree.SubElement(fill, qn("a:srgbClr"), {"val": hexa})
    else:
        etree.SubElement(ln, qn("a:noFill"))
    # The schema orders lnL, lnR, lnT, lnB before the cell fill.
    order = ["a:lnL", "a:lnR", "a:lnT", "a:lnB"]
    index = 0
    for child in list(tc_pr):
        if child.tag in [qn(t) for t in order[: order.index(tag)]]:
            index = list(tc_pr).index(child) + 1
    tc_pr.insert(index, ln)


def _table(shapes, item: dict, geo: Geometry, fonts: set[str]) -> None:
    from pptx.enum.text import MSO_ANCHOR  # ty: ignore[unresolved-import]
    from pptx.oxml.ns import qn  # ty: ignore[unresolved-import]
    from pptx.util import Emu  # ty: ignore[unresolved-import]

    rows = item["rows"]
    # Lay cells on a grid, as HTML's table model does with row and column spans.
    grid: dict[tuple[int, int], dict] = {}
    placed = []
    for ri, row in enumerate(rows):
        ci = 0
        for cell in row["cells"]:
            while (ri, ci) in grid:
                ci += 1
            for dr in range(cell["rowspan"]):
                for dc in range(cell["colspan"]):
                    grid[(ri + dr, ci + dc)] = cell
            placed.append((ri, ci, cell))
            ci += cell["colspan"]
    n_rows = max(r for r, _ in grid) + 1
    n_cols = max(c for _, c in grid) + 1
    shape = shapes.add_table(
        n_rows,
        n_cols,
        geo.emu(item["x"]),
        geo.emu(item["y"]),
        geo.emu(item["w"]),
        geo.emu(item["h"]),
    )
    table = shape.table
    table.first_row = table.horz_banding = False
    style = table._tbl.tblPr.find(qn("a:tableStyleId"))
    if style is not None:
        style.text = NO_TABLE_STYLE
    widths = [0.0] * n_cols
    for _, ci, cell in placed:
        if cell["colspan"] == 1:
            widths[ci] = max(widths[ci], cell["w"])
    spare = item["w"] - sum(widths)
    empty = [c for c, w in enumerate(widths) if not w]
    for c in empty:
        widths[c] = spare / len(empty) if spare > 0 else item["w"] / n_cols
    # The same slack as a wrapped text box, so a cell that fit on one line in Chromium still
    # does where the font sets a little wider; never past the slide's right edge.
    room = max(item["w"], geo.width * TABLE_RIGHT - item["x"])
    grow = min(1 + SLACK_SHARE, room / sum(widths)) if sum(widths) else 1
    widths = [w * max(grow, 1) for w in widths]
    shape.width = geo.emu(sum(widths))
    for c, w in enumerate(widths):
        table.columns[c].width = geo.emu(w)
    for r in range(n_rows):
        height = rows[r]["h"] if r < len(rows) else item["h"] / n_rows
        table.rows[r].height = geo.emu(height)
    valign = {"middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}
    for ri, ci, cell in placed:
        target = table.cell(ri, ci)
        if cell["rowspan"] > 1 or cell["colspan"] > 1:
            corner = (min(n_rows, ri + cell["rowspan"]) - 1, min(n_cols, ci + cell["colspan"]) - 1)
            try:
                target.merge(table.cell(*corner))
            except ValueError:
                # HTML lets spans overlap; PowerPoint does not. The cell stays unmerged.
                pass
        top, right, bottom, left = (geo.emu(v) for v in cell["pad"])
        target.margin_top, target.margin_right = Emu(top), Emu(right)
        target.margin_bottom, target.margin_left = Emu(bottom), Emu(left)
        target.vertical_anchor = valign.get(cell["valign"], MSO_ANCHOR.TOP)
        if cell["bg"][3] > 0:
            target.fill.solid()
            target.fill.fore_color.rgb = _rgb(cell["bg"])
        else:
            target.fill.background()
        target.text_frame.word_wrap = True
        _fill_text(target.text_frame, cell["runs"], geo.pt_per_px, cell, fonts)
        for side, border in cell["borders"].items():
            _cell_border(target, side, border, geo)


@dataclass
class Geometry:
    """Screen pixels of the browser's frame to EMU and points."""

    width: float  # frame width in screen px
    height: float

    @property
    def emu_per_px(self) -> float:
        return SLIDE_W_EMU / self.width

    @property
    def pt_per_px(self) -> float:
        return SLIDE_W_PT / self.width

    @property
    def px_per_slide_px(self) -> float:
        """Screen px per px of a 1600-wide slide, for constants written in slide px."""
        return self.width / 1600

    def emu(self, px: float) -> int:
        return round(px * self.emu_per_px)


def build(slides: list[Captured], title_only_layout: int, alt_text) -> tuple[Any, dict]:
    """The Presentation, and a tally of what it holds."""
    from pptx import Presentation  # ty: ignore[unresolved-import]
    from pptx.util import Emu  # ty: ignore[unresolved-import]

    prs = Presentation()
    first = slides[0].frame
    prs.slide_width = Emu(SLIDE_W_EMU)
    prs.slide_height = Emu(round(SLIDE_W_EMU * first["height"] / first["width"]))
    layout = prs.slide_layouts[title_only_layout]
    tally = {"text": 0, "table": 0, "picture": 0, "solid": 0, "backdrop": 0, "fonts": set(),
             "cut": []}  # fmt: skip
    aspect = first["height"] / first["width"]
    for number, cap in enumerate(slides, 1):
        geo = Geometry(cap.frame["width"], cap.frame["height"])
        slide = prs.slides.add_slide(layout)
        colour = uniform_colour(cap.background)
        if colour is not None:
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = _rgb(colour)
            tally["solid"] += 1
        else:
            _set_background_picture(slide, cap.background)
            tally["backdrop"] += 1
        if abs(cap.frame["height"] / cap.frame["width"] - aspect) > 0.01:
            tally["cut"].append(number)
        heading = slide.shapes.title
        titled = False
        # Pictures first, then tables, then text, so editable text is never under a picture;
        # within each kind, document order.
        rank = {"figure": 0, "equation": 0, "table": 1, "text": 2}
        for item in sorted(cap.items, key=lambda it: rank[it["kind"]]):
            if item["kind"] == "text":
                holder = heading if item.get("title") and not titled else None
                shape = _text_box(slide.shapes, item, geo, tally["fonts"], holder)
                if holder is not None:  # the layout's placeholder moves up to its turn
                    tree = shape._element.getparent()
                    tree.remove(shape._element)
                    tree.append(shape._element)
                titled = titled or holder is not None
                shape.name = "Title" if holder is not None else f"Text {tally['text'] + 1}"
                tally["text"] += 1
            elif item["kind"] == "table":
                _table(slide.shapes, item, geo, tally["fonts"])
                tally["table"] += 1
            elif item["shot"] in cap.shots:
                pic = slide.shapes.add_picture(
                    io.BytesIO(cap.shots[item["shot"]]),
                    geo.emu(item["x"]),
                    geo.emu(item["y"]),
                    geo.emu(item["w"]),
                    geo.emu(item["h"]),
                )
                pic.name = "Equation" if item["kind"] == "equation" else "Figure"
                pic._element.nvPicPr.cNvPr.set("descr", item["alt"] or pic.name)
                tally["picture"] += 1
        if not titled:
            # No heading: the title stays above the slide, read by the outline and screen
            # readers, as the picture export does.
            words = " ".join(
                r["text"] for it in cap.items if it["kind"] == "text" for r in it["runs"]
            )
            heading.text = alt_text(words, number)
            heading.left, heading.top = Emu(0), Emu(-int(prs.slide_height) // 6)
            heading.width, heading.height = prs.slide_width, Emu(int(prs.slide_height) // 6)
        if cap.notes:
            slide.notes_slide.notes_text_frame.text = cap.notes
    return prs, tally
