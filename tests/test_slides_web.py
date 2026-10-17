"""graphe:slides on a reveal.js web deck, checked and rendered in a real Chromium.

Each case copies the starter deck and breaks one thing. Chromium comes from Playwright
(`playwright install chromium`); a machine without it declares so with
TECHNE_NO_BROWSER=1, or the suite fails rather than skipping every web case unseen.
"""

from __future__ import annotations

import importlib.util
import os
import pathlib
import re
import shutil
import subprocess

import pytest

ROOT = pathlib.Path(__file__).resolve().parent.parent
STARTER = ROOT / "plugins" / "graphe" / "skills" / "slides" / "templates" / "web"
OPTED_OUT = os.environ.get("TECHNE_NO_BROWSER") == "1"
HAVE_LIBS = all(
    importlib.util.find_spec(m) is not None for m in ("playwright", "axe_playwright_python")
)


def _chromium_runs() -> bool:
    if not HAVE_LIBS:
        return False
    from playwright.sync_api import sync_playwright  # ty: ignore[unresolved-import]

    try:
        with sync_playwright() as pw:
            pw.chromium.launch().close()
    except Exception:
        return False
    return True


HAVE_BROWSER = not OPTED_OUT and _chromium_runs()
needs_browser = pytest.mark.skipif(
    not HAVE_BROWSER, reason="no Chromium (declared by TECHNE_NO_BROWSER)"
)


def test_the_browser_is_present_or_its_absence_is_declared():
    """Green by absence is the failure this guards."""
    assert HAVE_BROWSER or OPTED_OUT, (
        "Playwright's Chromium is not available and TECHNE_NO_BROWSER is unset: every web "
        "deck case would have skipped. Run `make test-unit` after `uv run --with "
        "playwright==1.63.0 playwright install chromium`, or set TECHNE_NO_BROWSER=1."
    )


def test_the_pins_agree_everywhere_they_are_written(sl):
    """Every file writing a Playwright, axe or python-pptx version writes the one slides.py uses."""
    files = (
        "Makefile",
        ".github/workflows/validate.yml",
        "plugins/graphe/skills/slides/SKILL.md",
        "docs/skills/slides.md",
        ".claude/skill-context.md",
    )
    for path in files:
        text = (ROOT / path).read_text()
        for pin in re.findall(r"(?:playwright|axe-playwright-python|python-pptx)==[\w.]+", text):
            assert pin in (sl.PLAYWRIGHT, sl.AXE, sl.PPTX), f"{path} pins {pin}"
    for path in files[:4]:
        assert sl.PLAYWRIGHT in (ROOT / path).read_text(), f"{path} does not pin {sl.PLAYWRIGHT}"


def test_the_vendored_reveal_is_the_version_the_skill_names():
    head = (STARTER / "vendor" / "reveal" / "reveal.js").read_text()[:400]
    skill = (STARTER.parent.parent / "SKILL.md").read_text()
    assert "reveal.js" in head
    version = re.search(r"reveal\.js (\d+\.\d+\.\d+)", skill)
    assert version, "SKILL.md names no reveal.js version"
    package = STARTER / "vendor" / "reveal" / "VERSION"
    assert package.read_text().strip() == version.group(1)
    assert f"`{version.group(1)}`" in (STARTER / "vendor" / "reveal" / "reveal.js").read_text()


@pytest.fixture
def web(tmp_path):
    """A copy of the starter deck, with a helper that edits one of its files."""
    deck = tmp_path / "deck"
    shutil.copytree(STARTER, deck)

    def edit(old: str, new: str, name: str = "index.html") -> pathlib.Path:
        path = deck / name
        text = path.read_text()
        assert old in text, f"{old!r} not in {name}"
        path.write_text(text.replace(old, new, 1))
        return deck / "index.html"

    edit.index = deck / "index.html"  # ty: ignore[unresolved-attribute]
    return edit


CLAIM = "<h2>[Claim headline: one full sentence the figure proves]</h2>"


def _check(sl, path, **kw):
    return sl.check(path, **kw)


@needs_browser
def test_the_starter_passes_with_nothing_to_review(sl, web):
    found = _check(sl, web.index)
    assert [str(f) for f in found if f.severity != "INFO"] == []
    assert sl.verdict(found)[0] == 0


@needs_browser
def test_a_slide_without_a_heading_blocks(sl, web, gates):
    found = _check(sl, web(CLAIM, "<p>[No heading]</p>"))
    assert ("BLOCK", "no-title") in gates(found)


@needs_browser
def test_low_contrast_text_blocks(sl, web, gates):
    found = _check(sl, web(CLAIM, CLAIM + '<p style="color:#9ca3af">faint</p>'))
    assert ("BLOCK", "contrast") in gates(found)


@needs_browser
def test_an_image_without_alt_text_blocks(sl, web, gates):
    pixel = "data:image/gif;base64,R0lGODlhAQABAAAAACw="
    found = _check(sl, web(CLAIM, CLAIM + f'<img src="{pixel}" width="40" height="40">'))
    assert ("BLOCK", "alt-text") in gates(found)


@needs_browser
def test_a_missing_file_blocks(sl, web, gates):
    found = _check(sl, web(CLAIM, CLAIM + '<img src="figs/gone.png" alt="gone">'))
    assert ("BLOCK", "asset") in gates(found)


@needs_browser
def test_a_value_with_no_data_behind_it_blocks(sl, web, gates):
    found = _check(sl, web(CLAIM, CLAIM + '<p data-value="results.nothing"></p>'))
    assert ("BLOCK", "script") in gates(found)


@needs_browser
def test_an_em_dash_blocks(sl, web, gates):
    found = _check(sl, web(CLAIM, CLAIM + "<p>before — after</p>"))
    assert ("BLOCK", "em-dash") in gates(found)


@needs_browser
def test_a_network_dependency_warns(sl, web, gates):
    link = '<link rel="stylesheet" href="https://example.invalid/font.css">'
    found = _check(sl, web('<link rel="stylesheet" href="deck.css">', link))
    assert ("WARN", "offline") in gates(found)


@needs_browser
def test_data_no_slide_binds_warns_with_its_path(sl, web, gates):
    """A figure left in the data file is published with the deck though no slide shows it."""
    web("    total: 1761,", "    total: 1761,\n    withheld: { early: 0.4 },", "data/deck-data.js")
    found = _check(sl, web.index)
    assert ("WARN", "unbound-data") in gates(found)
    (finding,) = [f for f in found if f.gate == "unbound-data"]
    assert "results.withheld" in finding.message
    assert "results.failShare" not in finding.message


@needs_browser
def test_a_binding_the_deck_drops_still_counts(sl, web, gates):
    """A data-if block that is false leaves the page, but its bindings are still the deck's."""
    web("    total: 1761,", "    total: 1761,\n    withheld: { early: 0.4 },", "data/deck-data.js")
    gone = '<p id="gone" data-source="results.withheld"></p><script>gone.remove()</script>'
    found = _check(sl, web(CLAIM, CLAIM + gone))
    assert ("WARN", "unbound-data") not in gates(found)


@needs_browser
def test_a_hidden_slide_does_not_bind_its_data(sl, web, gates):
    """reveal.js removes a data-visibility="hidden" slide, so its numbers are shown nowhere."""
    web("    total: 1761,", "    total: 1761,\n    withheld: { early: 0.4 },", "data/deck-data.js")
    cut = (
        '<section data-visibility="hidden"><h2>Cut</h2>'
        '<p data-value="results.withheld.early"></p></section>'
    )
    found = _check(sl, web("</section>", "</section>\n" + cut))
    assert ("WARN", "unbound-data") in gates(found)


@needs_browser
def test_a_reveal_attribute_naming_a_key_is_not_a_binding(sl, web, gates):
    web("    total: 1761,", "    total: 1761,\n    withheld: { early: 0.4 },", "data/deck-data.js")
    found = _check(sl, web("<section", '<section data-state="results"', "index.html"))
    assert ("WARN", "unbound-data") in gates(found)


@needs_browser
def test_a_key_outside_ascii_resolves_as_the_deck_reads_it(sl, web, gates):
    web("    total: 1761,", '    total: 1761,\n    "taux échec": 0.4,', "data/deck-data.js")
    found = _check(sl, web(CLAIM, CLAIM + '<p data-value="results.taux échec"></p>'))
    assert ("WARN", "unbound-data") not in gates(found)


@needs_browser
def test_data_that_throws_when_read_leaves_the_check_standing(sl, web, gates):
    getter = "    total: 1761,\n    get broken() { throw new Error('boom'); },"
    web("    total: 1761,", getter, "data/deck-data.js")
    found = _check(sl, web.index)
    assert ("WARN", "unbound-data") in gates(found)
    assert all(f.severity != "ERROR" for f in found), [str(f) for f in found]


@needs_browser
def test_binding_a_parent_binds_everything_under_it(sl, web, gates):
    web("    total: 1761,", "    total: 1761,\n    withheld: { early: 0.4 },", "data/deck-data.js")
    found = _check(sl, web(CLAIM, CLAIM + '<div data-source="results.withheld" hidden></div>'))
    assert ("WARN", "unbound-data") not in gates(found)


@needs_browser
def test_text_past_the_slide_edge_warns(sl, web, gates):
    wide = '<p style="width:2400px">[a line far wider than the slide]</p>'
    found = _check(sl, web(CLAIM, CLAIM + wide))
    assert ("WARN", "overflow") in gates(found)


@needs_browser
def test_content_past_the_edge_on_a_phone_blocks_at_both_widths(sl, web, gates):
    wide = '<div style="width:900px;height:20px;background:#000">[wider than a phone]</div>'
    found = _check(sl, web(CLAIM, CLAIM + wide))
    assert ("BLOCK", "phone-clip") in gates(found)
    clip = [f.message for f in found if f.gate == "phone-clip"]
    assert any("390px" in m for m in clip) and any("360px" in m for m in clip)


@needs_browser
def test_small_text_warns_on_desktop_and_phone(sl, web, gates):
    found = _check(sl, web(CLAIM, CLAIM + '<p style="font-size:11px">[fine print]</p>'))
    assert {("WARN", "small-text"), ("WARN", "phone")} <= gates(found)


@needs_browser
def test_motion_that_ignores_reduced_motion_warns(sl, web, gates):
    css = (STARTER / "deck.css").read_text()
    reduced = css[css.index("@media (prefers-reduced-motion") : css.index("@media print")]
    spin = "h2 { animation: spin 3s linear infinite; }\n@keyframes spin { to { opacity: 0.9; } }\n"
    web(reduced, spin, "deck.css")
    found = _check(sl, web.index)
    assert ("WARN", "motion") in gates(found)


@needs_browser
def test_a_talk_slide_without_notes_warns(sl, web, gates):
    note = '<aside class="notes">[Here\'s the problem.'
    path = web(note, '<aside class="unused">[Here\'s the problem.')
    found = _check(sl, path)
    assert ("WARN", "no-notes") in gates(found)


@needs_browser
def test_a_page_that_is_not_a_reveal_deck_is_unreadable(sl, tmp_path, gates):
    page = tmp_path / "page.html"
    page.write_text("<!doctype html><title>t</title><h1>Hello</h1>")
    assert ("ERROR", "unreadable") in gates(_check(sl, page))


@needs_browser
def test_render_writes_slides_phone_views_and_a_print_pdf(sl, web, tmp_path):
    out = tmp_path / "render"
    assert sl.render(web.index, out) == 0
    count = len(sl.html_slides(web.index))
    assert len(sorted(out.glob("slide-*.png"))) == count
    assert len(sorted(out.glob("phone-[0-9]*.png"))) == count
    pdf = out / "index.pdf"
    assert pdf.read_bytes()[:5] == b"%PDF-"
    if shutil.which("pdfinfo"):
        info = subprocess.run(["pdfinfo", str(pdf)], capture_output=True, text=True).stdout
        assert re.search(rf"Pages:\s+{count}\b", info)


@needs_browser
def test_pptx_exports_every_slide_as_a_picture_with_its_script(sl, web, tmp_path):
    from pptx import Presentation  # ty: ignore[unresolved-import]
    from pptx.enum.shapes import MSO_SHAPE_TYPE  # ty: ignore[unresolved-import]

    out = tmp_path / "deck.pptx"
    assert sl.pptx_html(web.index, out, tmp_path / "images") == 0
    expected = sl.html_slides(web.index)
    prs = Presentation(str(out))
    assert (prs.slide_width, prs.slide_height) == sl.SLIDE_EMU
    assert len(prs.slides) == len(expected)
    for number, (slide, (title, notes)) in enumerate(zip(prs.slides, expected, strict=True), 1):
        pictures = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
        assert len(pictures) == 1
        frame = (pictures[0].left, pictures[0].top, pictures[0].width, pictures[0].height)
        assert frame == (0, 0, *sl.SLIDE_EMU)
        heading = slide.shapes.title
        assert heading.text == (title or f"Slide {number}")
        assert heading.top + heading.height <= 0, "the title must sit off the slide"
        assert slide.notes_slide.notes_text_frame.text == notes
    assert len(list((tmp_path / "images").glob("slide-*.png"))) == len(expected)
    assert all("\n" not in s.shapes.title.text for s in prs.slides)
    for slide in prs.slides:
        pic = next(s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE)
        alt = pic._element.nvPicPr.cNvPr.get("descr")
        assert alt.startswith(slide.shapes.title.text.rstrip(".")[:20]), alt
        assert "\n" not in alt and len(alt) <= sl.ALT_MAX
    # The export passes the .pptx gates it can answer: a title and alt text on every slide.
    blocks = [f for f in sl.check(out) if f.severity == "BLOCK"]
    assert blocks == []


@pytest.mark.parametrize("name", ["export.pdf", "export"])
def test_pptx_writes_only_a_pptx(sl, web, tmp_path, name):
    with pytest.raises(SystemExit, match=r"must be a \.pptx"):
        sl.pptx_html(web.index, tmp_path / name)
    assert not (tmp_path / name).exists()


@pytest.mark.parametrize("where", ["same", "inside"])
def test_pptx_images_cannot_hold_the_pptx(sl, web, tmp_path, where):
    images = tmp_path / "out.pptx" if where == "same" else tmp_path / "images"
    target = images if where == "same" else images / "deck.pptx"
    with pytest.raises(SystemExit, match="cannot hold"):
        sl.pptx_html(web.index, target, images)
    assert not images.exists()


def test_pptx_refuses_a_target_under_a_file(sl, web, tmp_path):
    (tmp_path / "file").write_text("mine")
    with pytest.raises(SystemExit, match="is a file"):
        sl.pptx_html(web.index, tmp_path / "file" / "deck.pptx")
    assert (tmp_path / "file").read_text() == "mine"


@needs_browser
def test_pptx_images_refuse_a_folder_render_did_not_make(sl, web, tmp_path):
    images = tmp_path / "mine"
    images.mkdir()
    (images / "keep.txt").write_text("mine")
    with pytest.raises(SystemExit, match="not empty"):
        sl.pptx_html(web.index, tmp_path / "deck.pptx", images)
    assert (images / "keep.txt").read_text() == "mine"
    assert not (tmp_path / "deck.pptx").exists()


# ------------------------------------------------- cases the first review found --

PIXEL = "data:image/gif;base64,R0lGODlhAQABAIAAAAAAAP///yH5BAEAAAAALAAAAAABAAEAAAICRAEAOw=="


@needs_browser
def test_a_slide_over_an_image_is_counted_unchecked_not_judged(sl, web, gates):
    """axe cannot see an image behind text; the slide must not pass or block on a guess."""
    section = (
        '<section data-chapter="1 · The problem">\n'
        '  <span class="kicker">1 · The problem</span>\n  <h2>[One concrete'
    )
    pictured = section.replace(
        "<section ", f'<section data-background-image="{PIXEL}" style="color:#ffffff" ', 1
    )
    found = _check(sl, web(section, pictured))
    assert ("BLOCK", "contrast") not in gates(found)
    unchecked = [f for f in found if f.gate == "contrast" and f.severity == "INFO"]
    assert unchecked and int(unchecked[0].message.split()[0]) > 29


@needs_browser
def test_aa_accepts_what_aaa_blocks_and_aaa_names_its_own_threshold(sl, web, gates):
    path = web(CLAIM, CLAIM + '<p style="color:#6b6b6b;font-size:20px">[grey line]</p>')
    assert ("BLOCK", "contrast") not in gates(_check(sl, path, level="AA"))
    blocks = [f for f in _check(sl, path) if f.gate == "contrast" and f.severity == "BLOCK"]
    assert blocks and "needs 7:1 (AAA)" in blocks[0].message


@needs_browser
def test_notes_in_data_notes_count(sl, web, gates):
    path = web(
        '<section data-chapter="2 · How it works">',
        '<section data-chapter="2 · How it works" data-notes="Spoken script here.">',
    )
    web(
        '<aside class="notes">[One sentence per step',
        '<aside class="unused">[One sentence per step',
    )
    assert ("WARN", "no-notes") not in gates(_check(sl, path))
    assert "Spoken script here." in sl.script(path)[1]


@needs_browser
def test_a_heading_inside_the_notes_is_not_the_title(sl, web, gates):
    path = web(CLAIM, '<aside class="notes"><h2>Only in the notes</h2></aside>')
    assert ("BLOCK", "no-title") in gates(_check(sl, path))


@needs_browser
def test_overflow_is_measured_against_the_canvas_on_a_stock_deck(sl, tmp_path, gates):
    """A stock deck sizes each section to its content, so the section is no frame."""
    shutil.copytree(STARTER / "vendor", tmp_path / "vendor")
    lines = "".join(f"<p>[line {n}]</p>" for n in range(30))
    deck = tmp_path / "stock.html"
    deck.write_text(
        '<!doctype html><html lang="en"><head><meta charset="utf-8">'
        '<meta name="viewport" content="width=device-width, initial-scale=1"><title>t</title>'
        '<link rel="stylesheet" href="vendor/reveal/reveal.css"></head><body>'
        '<div class="reveal"><div class="slides">'
        f'<section><h2>Opening</h2><aside class="notes">Hi.</aside></section>'
        f'<section><h2>Too long</h2>{lines}<aside class="notes">Read.</aside></section>'
        '</div></div><script src="vendor/reveal/reveal.js"></script>'
        "<script>Reveal.initialize({transition: 'none'});</script></body></html>"
    )
    found = _check(sl, deck)
    assert ("WARN", "overflow") in gates(found)
    assert any(f.gate == "overflow" and f.slide == 2 for f in found)


@needs_browser
def test_content_clipped_off_the_slide_is_still_reported(sl, web, gates):
    clipped = '<div style="overflow:hidden"><p style="width:2400px">[cut off]</p></div>'
    assert ("WARN", "overflow") in gates(_check(sl, web(CLAIM, CLAIM + clipped)))


@needs_browser
def test_motion_names_only_the_slide_that_moves(sl, web):
    css = (STARTER / "deck.css").read_text()
    reduced = css[css.index("@media (prefers-reduced-motion") : css.index("@media print")]
    spin = ".discussion h2 { animation: spin 3s linear infinite; }\n"
    web(reduced, spin + "@keyframes spin { to { opacity: 0.9; } }\n", "deck.css")
    motion = [f for f in _check(sl, web.index) if f.gate == "motion"]
    assert len(motion) == 1
    named = {int(n) for n in re.findall(r"\d+", motion[0].message.split("[")[1].split("]")[0])}
    # The spinning slide is named. Slides with nothing that moves are not; with the rule gone,
    # the starter's dots (5) and timeline (8) animate too, and a fragment's fade may still be
    # running when it is read.
    assert 9 in named and not named & {1, 2, 4, 10, 11}


@needs_browser
def test_a_deck_that_never_starts_says_why(sl, web, gates):
    path = web(
        '(function () {\n  "use strict";',
        '(function () {\n  throw new Error("boom-before-init");',
        "deck.js",
    )
    found = _check(sl, path)
    assert ("ERROR", "unreadable") in gates(found)
    assert "boom-before-init" in found[0].message


@needs_browser
def test_empty_related_work_leaves_the_deck_running(sl, web, gates):
    data = (STARTER / "data" / "deck-data.js").read_text()
    works = data[data.index("    works: [") : data.index("    ],\n  },\n};")]
    found = _check(sl, web(works, "    works: [\n", "data/deck-data.js"))
    assert ("ERROR", "unreadable") not in gates(found)
    assert any(f.gate == "script" and "no works to draw" in f.message for f in found)


@needs_browser
def test_a_deck_with_no_slides_is_unreadable(sl, web, gates):
    text = web.index.read_text()
    body = text[text.index('<div class="slides">') + 20 : text.rindex("</div>\n</div>")]
    found = _check(sl, web(body, "\n"))
    assert ("ERROR", "unreadable") in gates(found)


@needs_browser
def test_a_missing_viewport_tag_is_a_phone_warning(sl, web):
    path = web('<meta name="viewport" content="width=device-width, initial-scale=1">', "")
    assert any(f.gate == "phone" and "viewport" in f.message for f in _check(sl, path))


@needs_browser
def test_figures_jargon_and_backup_from_work_on_a_web_deck(sl, web):
    path = web(CLAIM, CLAIM + "<p>[Gerrit] stopped 42% of them</p>")
    by = {(f.gate, f.slide) for f in _check(sl, path, jargon=("Gerrit",))}
    assert ("figures", 3) in by and ("jargon", 3) in by
    later = {(f.gate, f.slide) for f in _check(sl, path, jargon=("Gerrit",), backup_from=3)}
    assert not {("figures", 3), ("jargon", 3)} & later


def _stock(tmp_path, slides: str, config: str = "transition: 'none'") -> pathlib.Path:
    """A reveal.js deck with no starter CSS or JS: the vendored files and the given slides."""
    if not (tmp_path / "vendor").exists():
        shutil.copytree(STARTER / "vendor", tmp_path / "vendor")
    deck = tmp_path / "stock.html"
    deck.write_text(
        '<!doctype html><html lang="en"><head><meta charset="utf-8">'
        '<meta name="viewport" content="width=device-width, initial-scale=1"><title>t</title>'
        '<link rel="stylesheet" href="vendor/reveal/reveal.css"></head><body>'
        f'<div class="reveal"><div class="slides">{slides}</div></div>'
        '<script src="vendor/reveal/reveal.js"></script>'
        f"<script>Reveal.initialize({{{config}}});</script></body></html>"
    )
    return deck


OPENING = '<section><h2>Opening</h2><aside class="notes">Hi.</aside></section>'


@needs_browser
def test_a_section_with_its_own_background_is_measured_against_it(sl, web, gates):
    dark = '<section class="discussion dark" data-background-color="#14181f"'
    own = '<section class="discussion dark" style="background:#14181f"'
    assert ("BLOCK", "contrast") not in gates(_check(sl, web(dark, own)))
    grey = web('<p class="muted">[Example answers', '<p style="color:#3a3a3a">[Example answers')
    assert ("BLOCK", "contrast") in gates(_check(sl, grey))


@needs_browser
def test_a_gradient_on_the_section_itself_is_unchecked(sl, web, gates):
    dark = '<section class="discussion dark" data-background-color="#14181f"'
    painted = '<section class="discussion" style="background:linear-gradient(#000,#14181f)"'
    assert ("BLOCK", "contrast") not in gates(_check(sl, web(dark, painted)))


@needs_browser
def test_a_vertical_stack_keeps_phone_slides_aligned(sl, tmp_path):
    lines = "".join(f"<p>[line {n}]</p>" for n in range(30))
    deck = _stock(
        tmp_path,
        OPENING + '<section><section><h2>Top</h2><aside class="notes">a</aside></section>'
        f'<section><h2>Long</h2>{lines}<aside class="notes">b</aside></section></section>',
    )
    phone = [f for f in _check(sl, deck) if f.gate == "phone-clip"]
    assert phone and "slides [3]" in phone[0].message
    out = tmp_path / "render"
    sl.render(deck, out)
    assert len(list(out.glob("phone-[0-9]*.png"))) == 3


@needs_browser
def test_a_load_reveal_cancels_is_not_a_missing_file(sl, tmp_path, gates):
    (tmp_path / "frame.html").write_text("<!doctype html><title>f</title><p>frame")
    deck = _stock(
        tmp_path,
        OPENING + '<section data-background-iframe="frame.html"><h2>Framed</h2>'
        '<aside class="notes">x</aside></section>'
        '<section><h2>After</h2><aside class="notes">y</aside></section>',
    )
    assert ("BLOCK", "asset") not in gates(_check(sl, deck))


@needs_browser
def test_render_on_a_deck_with_no_slides_exits_cleanly(sl, tmp_path):
    out = tmp_path / "render"
    with pytest.raises(SystemExit, match="no slides"):
        sl.render(_stock(tmp_path, ""), out)
    assert not out.exists()


@needs_browser
def test_small_text_scales_with_the_deck_canvas(sl, tmp_path):
    big = _stock(
        tmp_path,
        OPENING + '<section><h2>Big canvas</h2><p style="font-size:21px">[fine print]</p>'
        '<aside class="notes">x</aside></section>',
        "transition: 'none', width: 1920, height: 1080",
    )
    small = [f for f in _check(sl, big) if f.gate == "small-text"]
    assert small and "1920x1080 canvas" in small[0].message


@needs_browser
def test_screen_reader_only_text_is_not_overflow(sl, web, gates):
    hidden = (
        '<span style="position:absolute;left:-10000px;width:1px;height:1px;overflow:hidden">'
        "[for screen readers]</span>"
    )
    found = _check(sl, web(CLAIM, CLAIM + hidden))
    assert ("WARN", "overflow") not in gates(found)


@needs_browser
def test_a_module_deck_is_told_to_set_window_reveal(sl, tmp_path):
    deck = _stock(tmp_path, OPENING)
    text = deck.read_text().replace(
        "<script>Reveal.initialize", "<script>const R = Reveal; Reveal = undefined; R.initialize"
    )
    deck.write_text(text)
    found = _check(sl, deck)
    assert found[0].gate == "unreadable" and "window.Reveal" in found[0].message


def test_the_pin_guard_reads_skill_templates_at_any_depth(tmp_path):
    guard = ROOT / "scripts" / "check_action_pins.sh"
    (tmp_path / "scripts").mkdir()
    shutil.copy(guard, tmp_path / "scripts")
    (tmp_path / ".github" / "workflows").mkdir(parents=True)
    shutil.copy(ROOT / ".github" / "workflows" / "docs.yml", tmp_path / ".github" / "workflows")
    run = lambda: subprocess.run(  # noqa: E731
        ["bash", str(tmp_path / "scripts" / "check_action_pins.sh")],
        capture_output=True,
        text=True,
        timeout=60,
    )
    assert run().returncode == 0, "no skill templates must pass"
    deep = tmp_path / "plugins" / "p" / "skills" / "s" / "templates" / "a" / "b"
    deep.mkdir(parents=True)
    template = (STARTER.parent / "pages.yml").read_text()
    (deep / "x.yaml").write_text(template)
    assert run().returncode == 0
    pin = re.search(r"deploy-pages@([0-9a-f]{40})", template)
    assert pin, "pages.yml no longer pins deploy-pages"
    sha = pin.group(1)
    (deep / "x.yaml").write_text(template.replace(sha, "1" * 40))
    done = run()
    assert done.returncode == 1 and "x.yaml" in done.stderr


@needs_browser
def test_render_on_a_page_that_is_not_a_deck_leaves_no_marker(sl, tmp_path):
    page = tmp_path / "page.html"
    page.write_text("<!doctype html><title>t</title><h1>Hello</h1>")
    out = tmp_path / "render"
    with pytest.raises(SystemExit, match=r"not a reveal\.js deck"):
        sl.render(page, out)
    assert not out.exists()


@needs_browser
def test_render_refuses_a_folder_it_did_not_make(sl, web, tmp_path):
    out = tmp_path / "theirs"
    out.mkdir()
    (out / "keep.png").write_bytes(b"x")
    with pytest.raises(SystemExit, match="not empty"):
        sl.render(web.index, out)
    assert (out / "keep.png").exists()


# ------------------------------------------------ cases the third review found --


@needs_browser
def test_a_see_through_section_is_blended_over_the_slide_colour(sl, tmp_path, gates):
    deck = _stock(
        tmp_path,
        OPENING + '<section data-background-color="#000000" '
        'style="background:rgba(255,255,255,0.2);color:#eeeeee"><h2>Over black</h2>'
        '<aside class="notes">x</aside></section>',
    )
    assert ("BLOCK", "contrast") not in gates(_check(sl, deck))


@needs_browser
def test_clipped_screen_reader_text_is_skipped_but_off_slide_text_is_not(sl, web, gates):
    hidden = (
        '<span style="position:absolute;width:1px;height:1px;overflow:hidden;'
        'clip:rect(0 0 0 0);white-space:nowrap"><strong>[a long phrase for screen readers'
        " only, wider than anything]</strong></span>"
    )
    assert ("WARN", "overflow") not in gates(_check(sl, web(CLAIM, CLAIM + hidden)))
    parked = '<p style="position:absolute;left:-400px;top:200px">[parked off the slide]</p>'
    assert ("WARN", "overflow") in gates(_check(sl, web(hidden, parked)))


@needs_browser
def test_a_deck_that_starts_only_on_wide_screens_names_the_phone(sl, tmp_path):
    deck = _stock(tmp_path, OPENING)
    deck.write_text(
        deck.read_text().replace(
            "<script>Reveal.initialize", "<script>if (innerWidth > 600) Reveal.initialize"
        )
    )
    found = _check(sl, deck)
    phone = [f for f in found if f.gate == "unreadable"]
    assert phone and "started on the desktop but not on a phone" in phone[0].message


# ------------------------------------------------ cases the fourth review found --


@needs_browser
def test_the_script_reads_the_starter_from_reveal(sl):
    code, out = sl.script(STARTER / "index.html")
    assert code == 0
    assert "## 3. [Claim headline: one full sentence the figure proves]" in out
    assert "Walk the figure left to right" in out
    _, backup = out.split("# Backup slides")
    assert "## 12. Backup slides" in backup


@needs_browser
def test_the_script_reads_notes_as_the_speaker_view_does(sl, tmp_path):
    deck = _stock(
        tmp_path,
        '<section data-notes="From the attribute."><h2>Line<br>break</h2>'
        '<aside class="notes">From the aside.</aside></section>'
        '<section><h2>Markup</h2><aside class="notes"><p>First sentence.</p><p>Second.</p>'
        "<ul><li>one</li><li>two</li></ul>Line<br>break</aside></section>"
        '<section data-visibility="hidden"><h2>Hidden</h2></section>'
        '<section><section><h2>Top</h2><aside class="notes">a</aside></section>'
        '<section><h2>Below</h2><aside class="notes">b</aside></section></section>',
    )
    slides = sl.html_slides(deck)
    assert [t for t, _ in slides] == ["Line break", "Markup", "Top", "Below"]
    assert slides[0][1] == "From the attribute."
    lines = [ln.strip() for ln in slides[1][1].splitlines() if ln.strip()]
    assert lines == ["First sentence.", "Second.", "one", "two", "Line", "break"]
    assert slides[3][1] == "b"


@needs_browser
def test_a_section_inside_slide_content_is_named(sl, tmp_path, gates):
    """reveal.js turns the slide into a vertical stack; the check says so instead of hiding it."""
    deck = _stock(
        tmp_path,
        OPENING + '<section><h2>Columns</h2><div><section class="col">left</section>'
        '<section class="col">right</section></div><aside class="notes">x</aside></section>',
    )
    found = _check(sl, deck)
    assert ("WARN", "nested-section") in gates(found)


@needs_browser
def test_a_clipped_shape_that_spills_is_still_overflow(sl, web, gates):
    card = '<div style="clip-path:inset(0 round 12px);width:1700px">[a card too wide]</div>'
    assert ("WARN", "overflow") in gates(_check(sl, web(CLAIM, CLAIM + card)))


@needs_browser
def test_a_phone_failure_is_not_blamed_on_files_missing_everywhere(sl, tmp_path):
    deck = _stock(tmp_path, OPENING.replace("</h2>", '</h2><img src="gone.png" alt="gone">'))
    deck.write_text(
        deck.read_text().replace(
            "<script>Reveal.initialize", "<script>if (innerWidth > 600) Reveal.initialize"
        )
    )
    phone = [f for f in _check(sl, deck) if f.gate == "unreadable"]
    assert phone and "started on the desktop but not on a phone" in phone[0].message
    assert "gone.png" not in phone[0].message


@needs_browser
def test_pptx_alt_text_carries_the_body_and_skips_web_only_text(sl, web, tmp_path):
    from pptx import Presentation  # ty: ignore[unresolved-import]
    from pptx.enum.shapes import MSO_SHAPE_TYPE  # ty: ignore[unresolved-import]

    index = web(
        CLAIM, CLAIM + '<p>Body sentence one.</p><p class="web-only">Hover for a summary.</p>'
    )
    out = tmp_path / "deck.pptx"
    sl.pptx_html(index, out)
    alts = [
        s._element.nvPicPr.cNvPr.get("descr")
        for slide in Presentation(str(out)).slides
        for s in slide.shapes
        if s.shape_type == MSO_SHAPE_TYPE.PICTURE
    ]
    hit = [a for a in alts if "Body sentence one." in a]
    assert len(hit) == 1 and hit[0].startswith("[Claim headline")
    assert not any("Hover for a summary" in a for a in alts)


def _pdf_text(pdf: pathlib.Path) -> str:
    return subprocess.run(["pdftotext", str(pdf), "-"], capture_output=True, text=True).stdout


needs_poppler = pytest.mark.skipif(not shutil.which("pdftotext"), reason="no poppler")


@needs_browser
@needs_poppler
def test_pdf_has_one_vector_page_per_slide_with_text_and_working_links(sl, web, tmp_path):
    index = web(
        CLAIM,
        CLAIM
        + '<p>Body sentence one.</p><p class="web-only">Hover for a summary.</p>'
        + '<p><a href="https://example.test/paper">the paper</a> <a href="#/agenda">agenda</a></p>',
    )
    web(
        '<section>\n  <span class="kicker">Today',
        '<section id="agenda">\n  <span class="kicker">Today',
    )
    out = tmp_path / "deck.pdf"
    assert sl.pdf_html(index, out) == 0
    data = out.read_bytes()
    assert data[:5] == b"%PDF-" and not out.with_name("deck.pdf.partial").exists()
    count = len(sl.html_slides(index))
    info = subprocess.run(["pdfinfo", str(out)], capture_output=True, text=True).stdout
    assert re.search(rf"Pages:\s+{count}\b", info)
    text = _pdf_text(out)
    assert "Body sentence one." in text and "Hover for a summary" not in text
    # A link to an external page and one to a slide by id, both live in the PDF.
    assert b"/URI (https://example.test/paper)" in data
    assert b"/Dest /agenda" in data


@needs_browser
@needs_poppler
def test_pdf_shows_each_slide_fully_built_whatever_the_deck_configures(sl, web, tmp_path):
    web("pdfSeparateFragments: false,", "", "deck.js")
    index = web(CLAIM, CLAIM + '<p class="fragment">first</p><p class="fragment">second</p>')
    out = tmp_path / "deck.pdf"
    sl.pdf_html(index, out)
    info = subprocess.run(["pdfinfo", str(out)], capture_output=True, text=True).stdout
    assert re.search(rf"Pages:\s+{len(sl.html_slides(index))}\b", info)
    assert "first" in _pdf_text(out) and "second" in _pdf_text(out)


@needs_browser
@needs_poppler
def test_pdf_keeps_text_pinned_to_the_bottom_of_a_slide(sl, web, tmp_path):
    css = web.index.parent / "deck.css"
    css.write_text(
        css.read_text()
        + ".reveal .slides section { height: 100%; } .footnote { position: absolute; bottom: 0; }"
    )
    out = tmp_path / "deck.pdf"
    sl.pdf_html(web.index, out)
    assert "Where the discussion stops fall" in _pdf_text(out)


@pytest.mark.parametrize("name", ["export.pptx", "export"])
def test_pdf_writes_only_a_pdf(sl, web, tmp_path, name):
    with pytest.raises(SystemExit, match=r"must be a \.pdf"):
        sl.pdf_html(web.index, tmp_path / name)
    assert not (tmp_path / name).exists()


def test_pdf_refuses_a_folder_and_a_target_under_a_file(sl, web, tmp_path):
    (tmp_path / "folder.pdf").mkdir()
    with pytest.raises(SystemExit, match="is a folder"):
        sl.pdf_html(web.index, tmp_path / "folder.pdf")
    (tmp_path / "file").write_text("mine")
    with pytest.raises(SystemExit, match="is a file"):
        sl.pdf_html(web.index, tmp_path / "file" / "deck.pdf")
    assert (tmp_path / "file").read_text() == "mine"


def test_pdf_exports_only_a_web_deck(sl, tmp_path):
    (tmp_path / "deck.pdf").write_bytes(b"%PDF-")
    with pytest.raises(SystemExit, match="web deck"):
        sl.pdf_html(tmp_path / "deck.pdf", tmp_path / "out.pdf")


def test_alt_text_is_cut_at_a_word(sl):
    alt = sl._alt_text("Title", "word " * 1000, 1)
    assert len(alt) <= sl.ALT_MAX and alt.endswith("word\u2026")
    assert sl._alt_text("", "", 3) == "Slide 3"


@needs_browser
@pytest.mark.parametrize(
    ("viewport_tag", "width", "sideways"),
    [(True, 1000, True), (True, 300, False), (False, 300, False), (False, 1500, False)],
)
def test_the_sideways_check_fires_only_when_a_phone_pans(sl, viewport_tag, width, sideways):
    """A page with a viewport tag and content wider than the phone pans; a page without one is
    zoomed out to fit, so it does not."""
    from playwright.sync_api import sync_playwright  # ty: ignore[unresolved-import]

    meta = "<meta name=viewport content='width=device-width,initial-scale=1'>" * viewport_tag
    page_html = f"<!doctype html><head>{meta}</head><body><div style='width:{width}px'>x</div>"
    with sync_playwright() as pw:
        browser = pw.chromium.launch()
        ctx = browser.new_context(viewport=sl.PHONE, is_mobile=True, has_touch=True)
        page = ctx.new_page()
        page.set_content(page_html)
        assert page.evaluate(sl._SIDEWAYS) is sideways
        browser.close()
