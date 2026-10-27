"""graphe:slides `pptx --editable`: a web deck exported with its text and tables editable.

The custom deck fixture is exported once in a real Chromium and read back with python-pptx.
The render comparison opens the export in PowerPoint or LibreOffice, whichever is here, and
compares each slide with the browser's picture of it; a machine with neither declares so with
TECHNE_NO_OFFICE=1, or the suite fails rather than skipping the comparison unseen.
"""

# ruff: noqa: RUF001  # expected Unicode from TeX
from __future__ import annotations

import io
import os
import pathlib
import types

import pytest
from test_slides_web import needs_browser, web  # noqa: F401  # web is a fixture

ROOT = pathlib.Path(__file__).resolve().parent.parent
CUSTOM = ROOT / "tests" / "fixtures" / "custom-deck" / "index.html"
NO_OFFICE = os.environ.get("TECHNE_NO_OFFICE") == "1"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
# The fixture is 1600 px wide, so one of its pixels is this many EMU and points.
EMU_PX, PT_PX = 12192000 / 1600, 960 / 1600


def _office(sl) -> bool:
    return sl.powerpoint_available() or bool(sl.libreoffice())


# ------------------------------------------------------------------ pure --


@pytest.mark.parametrize(
    ("source", "expected"),
    [
        (r"\alpha^{2} \le \beta_{i}", "α² ≤ βᵢ"),
        (r"\frac{1}{2\pi}", "1/(2π)"),
        (r"x_{ij}", "xᵢⱼ"),
        (r"x_{qw}", "x_(qw)"),
        (r"\theta = [m\ c\ k]^\top", "θ = [m c k]ᵀ"),
        (r"\sqrt{k/m}", "√(k/m)"),
        (r"\mathbf{z} \odot \mathbf{s}", "z ⊙ s"),
        (r"\frac{x_{i}}{n}", "xᵢ/n"),
        (r"m\ddot{x}", "mẍ"),
        (r"\left( a \right)", "( a )"),
    ],
)
def test_inline_tex_reads_as_unicode(sl, source, expected):
    assert sl.pptx_editable.tex(source) == expected


def test_runs_collapse_whitespace_and_join_one_style(sl):
    style: dict = {"size": 20, "color": [0, 0, 0, 1], "family": "Arial", "shift": 0}
    style |= {"bold": False, "italic": False, "underline": False, "strike": False}
    style |= {"href": None, "spacing": 0}
    runs = [
        {**style, "text": "\n   A  custom "},
        {**style, "text": " deck"},
        {**style, "bold": True, "text": " bold "},
        {**style, "text": "\n", "pre": True},
        {**style, "tex": r"\alpha", "text": ""},
    ]
    out = [(r["text"], r["bold"]) for r in sl.pptx_editable._collapse(runs)]
    # The browser keeps one space where runs meet, at the start of the later run.
    assert out == [("A custom deck", False), (" bold ", True), ("\nα", False)]


# ------------------------------------------------------------ the export --


@pytest.fixture(scope="module")
def export(sl, tmp_path_factory) -> types.SimpleNamespace:
    from pptx import Presentation  # ty: ignore[unresolved-import]

    out = tmp_path_factory.mktemp("editable")
    target, images = out / "deck.pptx", out / "images"
    assert sl.pptx_html(CUSTOM, target, images, editable=True) == 0
    return types.SimpleNamespace(path=target, images=images, prs=Presentation(str(target)))


def _shapes(slide) -> dict:
    return {s.text_frame.text: s for s in slide.shapes if s.has_text_frame and s.text_frame.text}


def _run(slide, text: str):
    for shape in slide.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                for r in p.runs:
                    if r.text == text:
                        return r
    raise AssertionError(f"no run {text!r}; runs: {[s.text for s in _shapes(slide)]}")


@needs_browser
def test_every_slide_exports_with_a_16_by_9_page(export):
    assert len(export.prs.slides) == 4
    assert (export.prs.slide_width, export.prs.slide_height) == (12192000, 6858000)


@needs_browser
def test_text_stays_editable_text_with_its_formatting(export):
    one = export.prs.slides[0]
    title = one.shapes.title
    assert title.text_frame.text == "Editable export fixture"
    assert title.text_frame.paragraphs[0].runs[0].font.size.pt == pytest.approx(72 * PT_PX, abs=0.5)
    assert _run(one, "bold words").font.bold is True
    assert _run(one, "italic words").font.italic is True
    assert _run(one, "a link").hyperlink.address == "https://example.org/paper"
    assert "42%" in _shapes(one)


@needs_browser
def test_text_boxes_sit_where_the_browser_laid_them_out(export):
    title = export.prs.slides[0].shapes.title
    eyebrow = _shapes(export.prs.slides[0])["FIXTURE KICKER"]
    assert title.left == pytest.approx(100 * EMU_PX, abs=EMU_PX)
    assert eyebrow.top == pytest.approx(80 * EMU_PX, abs=2 * EMU_PX)
    assert title.top > eyebrow.top + eyebrow.height - EMU_PX


@needs_browser
def test_letter_spacing_and_text_transform_carry_over(export):
    run = _run(export.prs.slides[0], "FIXTURE KICKER")
    assert int(run._r.get_or_add_rPr().get("spc")) == round(0.2 * 24 * PT_PX * 100)


@needs_browser
def test_a_block_displayed_inline_tag_and_flex_items_get_boxes_of_their_own(export):
    assert {"375", "solves"} <= set(_shapes(export.prs.slides[0]))
    assert {"Group 1 · Fixture", "03 / 03"} <= set(_shapes(export.prs.slides[2]))


@needs_browser
def test_lists_get_native_bullets_and_numbers(export):
    two = export.prs.slides[1]
    first = _shapes(two)["First point with α² ≤ βᵢ"].text_frame.paragraphs[0]._p.pPr
    assert first.find(f"{A}buChar").get("char") == "•"
    assert int(first.get("indent")) < 0 < int(first.get("marL"))
    numbers = [
        _shapes(two)[t].text_frame.paragraphs[0]._p.pPr.find(f"{A}buAutoNum")
        for t in ("Third, numbered", "Fourth, numbered")
    ]
    assert [(n.get("type"), n.get("startAt")) for n in numbers] == [
        ("arabicPeriod", "3"),
        ("arabicPeriod", "4"),
    ]
    assert "Second point, shown by a fragment" in _shapes(two)


@needs_browser
def test_tables_are_native_with_no_theme_style(sl, export):
    tables = [s for s in export.prs.slides[2].shapes if s.has_table]
    assert len(tables) == 1
    table = tables[0].table
    assert [[c.text for c in row.cells] for row in table.rows] == [
        ["Method", "Error"],
        ["Gradient descent", "0.031"],
        ["Least squares", "0.007"],
    ]
    style = table._tbl.tblPr.find(f"{A}tableStyleId").text
    assert style == sl.pptx_editable.NO_TABLE_STYLE
    assert str(table.cell(0, 0).fill.fore_color.rgb) == "0B5C63"
    border = table.cell(1, 0)._tc.tcPr.find(f"{A}lnB")
    assert border.find(f"{A}solidFill/{A}srgbClr").get("val") == "9AA5A8"


@needs_browser
def test_figures_and_display_equations_are_pictures_with_alt_text(export):
    def alts(slide):
        return [s._element.nvPicPr.cNvPr.get("descr") for s in slide.shapes if s.shape_type == 13]

    assert alts(export.prs.slides[1]) == ["A teal square figure"]
    assert alts(export.prs.slides[2]) == [r"m\ddot{x} + c\dot{x} + kx = F(t)"]


@needs_browser
def test_notes_carry_over_one_line_per_paragraph(export):
    notes = [s.notes_slide.notes_text_frame.text for s in export.prs.slides if s.has_notes_slide]
    assert notes == [
        "Open with the number.\nThen the plan.",
        "Walk the list, then point at the figure.",
    ]


def _background(slide):
    from PIL import Image  # ty: ignore[unresolved-import]

    bg = slide._element.find(f"{P}cSld/{P}bg")
    blip = bg.find(f".//{A}blip") if bg is not None else None
    if blip is None:
        return None
    rid = blip.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
    return Image.open(io.BytesIO(slide.part.related_part(rid).blob)).convert("RGB")


@needs_browser
def test_a_one_colour_background_is_a_solid_fill(export):
    one = export.prs.slides[0]
    assert _background(one) is None
    assert str(one.background.fill.fore_color.rgb) == "14181F"


@needs_browser
def test_the_background_keeps_decoration_but_no_glyphs_and_no_deck_controls(export):
    image = _background(export.prs.slides[1])
    assert image is not None, "the card should make slide 2's background a picture"
    scale = image.width / 1600
    card = image.getpixel((int(140 * scale), int(460 * scale)))
    assert card == (0xE8, 0xF3, 0xF4)

    def plain(box):
        crop = image.crop(tuple(int(v * scale) for v in box))
        return {c for _, c in crop.getcolors(1 << 16)} == {(255, 255, 255)}

    assert plain((100, 820, 300, 860)), "the footer's glyphs are baked into the background"
    assert plain((600, 860, 1000, 900)), "the deck's own controls are in the background"


@needs_browser
def test_the_starter_reveal_deck_exports_editable(sl, web, tmp_path):  # noqa: F811
    from pptx import Presentation  # ty: ignore[unresolved-import]

    target = tmp_path / "starter.pptx"
    assert sl.pptx_html(web.index, target, editable=True) == 0
    prs = Presentation(str(target))
    assert len(prs.slides) == len(sl.html_slides(web.index))
    first = prs.slides[0]
    assert first.shapes.title.text_frame.text.startswith("[Talk title")
    assert first.notes_slide.notes_text_frame.text.startswith("[Hi, I'm")


@needs_browser
def test_a_page_that_is_not_a_deck_is_refused(sl, tmp_path):
    page = tmp_path / "page.html"
    page.write_text("<!doctype html><title>t</title><h1>Hello</h1>")
    with pytest.raises(SystemExit, match=r"no <section class=\"slide\"> elements"):
        sl.pptx_html(page, tmp_path / "out.pptx", editable=True)


def _picture(shape):
    from PIL import Image  # ty: ignore[unresolved-import]

    return Image.open(io.BytesIO(shape.image.blob)).convert("RGB")


def _colours(image, box, scale):
    crop = image.crop(tuple(int(v * scale) for v in box))
    return {c for _, c in crop.getcolors(1 << 20)}


@needs_browser
def test_generated_text_is_a_run_and_a_counter_stays_visible_in_the_background(export):
    edge = export.prs.slides[3]
    centred = _shapes(edge)["§ Centred"]
    assert centred.left + centred.width / 2 == pytest.approx(750 * EMU_PX, abs=8 * EMU_PX)
    kicker = _shapes(edge)["→ Kicker with arrow"]
    assert kicker.left == pytest.approx(100 * EMU_PX, abs=2 * EMU_PX)
    assert str(_run(edge, "→ ").font.color.rgb) == "B00020"
    assert "Measure first" in _shapes(edge)
    image = _background(edge)
    # "Step 1: " is drawn by a CSS counter, which the page cannot hand over as text.
    assert (0xB0, 0x00, 0x20) in _colours(image, (100, 200, 200, 236), image.width / 1600)
    # A block with only generated text and whitespace keeps it in the background, visible.
    assert (0xB0, 0x00, 0x20) in _colours(image, (600, 140, 700, 176), image.width / 1600)
    # The arrow became text, so the background has no glyph where it was.
    assert _colours(image, (100, 140, 125, 176), image.width / 1600) == {(255, 255, 255)}


@needs_browser
def test_display_contents_and_screen_reader_text(export):
    edge = _shapes(export.prs.slides[3])
    assert "Inside contents" in edge
    assert "Visible words" in edge
    assert not any("hidden words" in text for text in edge)


@needs_browser
def test_an_icon_in_currentcolor_and_gradient_text_stay_in_the_background(export):
    edge = export.prs.slides[3]
    image = _background(edge)
    scale = image.width / 1600
    # The icon sits after "Icon "; its pixels are the text colour, not transparent.
    assert (0x1A, 0x1A, 0x1A) in _colours(image, (150, 380, 260, 420), scale)
    assert not any("Gradient" in text for text in _shapes(edge))
    assert len(_colours(image, (100, 440, 300, 476), scale)) > 10


@needs_browser
def test_wrapped_text_gets_slack_and_one_line_text_does_not_wrap(export):
    edge = _shapes(export.prs.slides[3])
    long = edge["A paragraph long enough that the browser wraps it onto a second line."]
    assert long.text_frame.word_wrap is True
    assert long.width > 520 * EMU_PX * 1.05
    assert edge["Half ink"].text_frame.word_wrap is False
    # One line set from several inline elements is still one line.
    lead = _shapes(export.prs.slides[0])["A custom deck: bold words, italic words and a link."]
    assert lead.text_frame.word_wrap is False


@needs_browser
def test_an_item_that_opens_with_a_paragraph_gets_its_bullet(export):
    item = _shapes(export.prs.slides[3])["Item in a paragraph"]
    assert item.text_frame.paragraphs[0]._p.pPr.find(f"{A}buChar") is not None


@needs_browser
def test_an_inside_marker_stays_where_the_text_starts_with_room_after_it(export):
    item = _shapes(export.prs.slides[3])["Inside bullet"]
    assert item.left == pytest.approx(600 * EMU_PX, abs=2 * EMU_PX)
    ppr = item.text_frame.paragraphs[0]._p.pPr
    assert int(ppr.get("marL")) > 0 and ppr.find(f"{A}buChar") is not None


@needs_browser
def test_right_to_left_and_translucent_text(export):
    edge = export.prs.slides[3]
    assert _shapes(edge)["שלום עולם"].text_frame.paragraphs[0]._p.pPr.get("rtl") == "1"
    alpha = _run(edge, "Half ink")._r.find(f".//{A}srgbClr/{A}alpha")
    assert alpha is not None and alpha.get("val") == "50000"


@needs_browser
def test_text_over_a_figure_is_on_top_and_not_in_its_picture(export):
    edge = export.prs.slides[3]
    shapes = list(edge.shapes)
    figure = next(s for s in shapes if s.shape_type == 13 and "caption" in
                  s._element.nvPicPr.cNvPr.get("descr"))  # fmt: skip
    caption = _shapes(edge)["Over the figure"]
    assert shapes.index(caption) > shapes.index(figure)
    image = _picture(figure)
    # Inside the square, under the caption: teal only, no white glyphs.
    scale = image.width / 420
    assert _colours(image, (100, 115, 200, 140), scale) == {(0x08, 0x7C, 0x83)}


@needs_browser
def test_overlapping_table_spans_export_without_a_crash(export):
    tables = [s for s in export.prs.slides[3].shapes if s.has_table]
    assert len(tables) == 1


@needs_browser
def test_tablecell_borders_come_before_the_fill_as_the_schema_orders_them(export):
    table = next(s for s in export.prs.slides[2].shapes if s.has_table).table
    order = [child.tag.split("}")[1] for child in table.cell(0, 0)._tc.tcPr]
    assert order[:5] == ["lnL", "lnR", "lnT", "lnB", "solidFill"]


@needs_browser
def test_a_slide_larger_than_the_window_is_captured_whole(sl, tmp_path):
    from pptx import Presentation  # ty: ignore[unresolved-import]

    deck = tmp_path / "tall.html"
    deck.write_text(
        "<!doctype html><title>t</title><style>body{margin:0}"
        ".slide{position:relative;width:1600px;height:1400px;background:#fff;font:32px Arial}"
        "p{position:absolute;margin:0}</style>"
        '<section class="slide active"><h1>Tall</h1>'
        '<p style="top:1300px;left:100px">Bottom line</p>'
        '<div style="position:absolute;left:0;top:1350px;width:1600px;height:50px;'
        'background:#c00000"></div></section>'
    )
    target = tmp_path / "tall.pptx"
    assert sl.pptx_html(deck, target, editable=True) == 0
    prs = Presentation(str(target))
    assert prs.slide_height == round(12192000 * 1400 / 1600)
    bottom = _shapes(prs.slides[0])["Bottom line"]
    assert bottom.top == pytest.approx(1300 * EMU_PX, abs=4 * EMU_PX)
    image = _background(prs.slides[0])
    assert image.height / image.width == pytest.approx(1400 / 1600, abs=0.01)
    assert image.getpixel((800, image.height - 5)) == (0xC0, 0x00, 0x00)


@needs_browser
def test_a_centred_slide_wider_than_the_window_is_read_after_the_window_grows(sl, tmp_path):
    from pptx import Presentation  # ty: ignore[unresolved-import]

    deck = tmp_path / "wide.html"
    deck.write_text(
        "<!doctype html><title>t</title><style>body{margin:0}"
        ".slide{position:absolute;left:50%;top:0;transform:translateX(-50%);width:2400px;"
        "height:900px;background:#fff;font:40px Arial}p{position:absolute;margin:0}</style>"
        '<section class="slide active"><p style="left:100px;top:100px">Wide left text</p>'
        '<div style="position:absolute;right:0;bottom:0;width:200px;height:100px;'
        'background:#c00000"></div></section>'
    )
    target = tmp_path / "wide.pptx"
    assert sl.pptx_html(deck, target, editable=True) == 0
    prs = Presentation(str(target))
    text = _shapes(prs.slides[0])["Wide left text"]
    assert text.left == pytest.approx(100 * 12192000 / 2400, abs=12192000 / 2400 * 2)
    image = _background(prs.slides[0])
    scale = image.width / 2400
    assert image.getpixel((int(2350 * scale), int(880 * scale))) == (0xC0, 0x00, 0x00)
    assert _colours(image, (100, 100, 400, 150), scale) == {(255, 255, 255)}


@needs_browser
def test_a_slide_that_never_fits_is_exported_as_far_as_it_shows_and_named(sl, tmp_path, capsys):
    deck = tmp_path / "grows.html"
    deck.write_text(
        "<!doctype html><title>t</title><style>body{margin:0}"
        ".slide{position:absolute;left:0;top:0;width:130vw;height:600px;"
        "background:linear-gradient(90deg,#ffffff,#000000)}</style>"
        '<section class="slide active"><p>Always wider</p></section>'
    )
    from pptx import Presentation  # ty: ignore[unresolved-import]

    target = tmp_path / "out.pptx"
    assert sl.pptx_html(deck, target, editable=True) == 0
    assert "slides 1 lie partly off screen" in capsys.readouterr().out
    prs = Presentation(str(target))
    image = _background(prs.slides[0])
    # The visible part keeps its place and the off-screen rest takes the edge colour (white
    # here); stretching the visible part instead would end in the gradient's dark grey.
    assert image.getpixel((image.width - 2, image.height // 2)) == (255, 255, 255)
    assert sum(image.getpixel((int(image.width * 0.7), image.height // 2))) < 3 * 120


@needs_browser
def test_a_tall_slide_in_a_fit_to_window_deck_is_captured_whole(sl, tmp_path):
    from pptx import Presentation  # ty: ignore[unresolved-import]

    deck = tmp_path / "fit.html"
    deck.write_text(
        "<!doctype html><title>t</title><style>html,body{margin:0;height:100%}"
        "#deck{position:absolute;left:50%;top:50%;width:1600px;height:900px;"
        "transform-origin:0 0}"
        ".slide{position:absolute;left:0;top:0;width:1600px;height:900px;display:none;"
        "background:#fff;font:40px Arial}.slide.active{display:block}"
        ".tall{height:1300px}p{position:absolute;margin:0}</style><div id=deck>"
        '<section class="slide active"><p style="left:100px;top:100px">First</p></section>'
        '<section class="slide tall"><p style="left:100px;top:1200px">Low text</p>'
        '<div style="position:absolute;left:0;top:1250px;width:1600px;height:50px;'
        'background:#c00000"></div></section></div><script>'
        "const d=document.getElementById('deck');const fit=()=>{"
        "const s=Math.min(innerWidth/1600,innerHeight/900);"
        "d.style.transform=`scale(${s}) translate(-50%,-50%)`;};"
        "addEventListener('resize',fit);fit();</script>"
    )
    target = tmp_path / "fit.pptx"
    assert sl.pptx_html(deck, target, editable=True) == 0
    tall = Presentation(str(target)).slides[1]
    image = _background(tall)
    assert image.getpixel((image.width // 2, image.height - 3)) == (0xC0, 0x00, 0x00)


# ---------------------------------------------------------- render check --


def test_an_office_renderer_is_present_or_its_absence_is_declared(sl):
    """Green by absence is the failure this guards."""
    assert NO_OFFICE or _office(sl), (
        "neither PowerPoint nor LibreOffice is here and TECHNE_NO_OFFICE is unset: the render "
        "comparison would have skipped. Install LibreOffice, or set TECHNE_NO_OFFICE=1."
    )


def _distance(a, b) -> float:
    """Mean absolute grey-level difference of two pictures at 320x180, from 0 to 1."""
    from PIL import ImageChops, ImageStat  # ty: ignore[unresolved-import]

    small = [im.convert("L").resize((320, 180)) for im in (a, b)]
    return ImageStat.Stat(ImageChops.difference(*small)).mean[0] / 255


@needs_browser
def test_each_slide_renders_in_office_as_the_browser_shows_it(sl, export, tmp_path):
    if NO_OFFICE or not _office(sl):
        pytest.skip("no PowerPoint or LibreOffice (declared by TECHNE_NO_OFFICE)")
    from PIL import Image  # ty: ignore[unresolved-import]

    out = tmp_path / "render"
    assert sl.render(export.path, out, "auto", 96) == 0
    office = [Image.open(p) for p in sorted(out.glob("slide-*.png"))]
    browser = [Image.open(p) for p in sorted(export.images.glob("slide-*.png"))]
    assert len(office) == len(browser) == 4
    for number, (o, b) in enumerate(zip(office, browser, strict=True), 1):
        same = _distance(o, b)
        other = _distance(o, browser[number % 4])
        assert same < 0.03, f"slide {number} differs from the browser by {same:.3f}"
        assert same < other / 3, f"slide {number} is as close to another slide ({other:.3f})"
